"""Exports of the extraction dataset: long-format CSV, JSON, XLSX, and a ZIP bundle with R and Python loader scripts.

The long format has one row per value component (for example the mean, SD, and n of an arm's outcome), which loads
directly into R or pandas and can be reshaped for analysis.
"""

import csv
import io
import json
import zipfile
from typing import Any

from openpyxl import Workbook

COLUMNS = [
    "study_id",
    "study",
    "registry_ids",
    "field_id",
    "field",
    "section",
    "field_type",
    "outcome",
    "timepoint",
    "arm",
    "component",
    "value",
    "unit",
    "not_reported",
    "source",
    "flags",
    "derivation_method",
    "evidence_span_ids",
    "cell_state",
]


def rows(dataset: dict[str, Any]) -> list[dict[str, Any]]:
    fields = {item["id"]: item for item in dataset["fields"]}
    out = []
    for study in dataset["studies"]:
        for cell in study["values"]:
            item = fields.get(cell["field_id"], {})
            base = {
                "study_id": study["study_id"],
                "study": study["label"],
                "registry_ids": ";".join(study["registry_ids"]),
                "field_id": cell["field_id"],
                "field": cell["field"],
                "section": item.get("section", ""),
                "field_type": item.get("field_type", ""),
                "outcome": item.get("outcome", ""),
                "timepoint": item.get("timepoint", ""),
                "arm": cell["arm"] or "",
                "unit": cell["unit"],
                "not_reported": cell["not_reported"],
                "source": cell["source"] or "",
                "flags": ";".join(cell["flags"]),
                "derivation_method": (cell["derivation"] or {}).get("method", ""),
                "evidence_span_ids": ";".join(str(i) for i in cell["span_ids"]),
                "cell_state": cell["state"],
            }
            value = cell["value"]
            if not value:
                out.append({**base, "component": "", "value": ""})
                continue
            for component, item_value in value.items():
                if isinstance(item_value, list):
                    item_value = ";".join(str(v) for v in item_value)
                out.append({**base, "component": component, "value": item_value})
    return out


def to_csv(dataset: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS)
    writer.writeheader()
    writer.writerows(rows(dataset))
    return buffer.getvalue().encode()


def to_json(dataset: dict[str, Any], metadata: dict[str, Any]) -> bytes:
    return json.dumps({"metadata": metadata, **dataset}, indent=2, default=str).encode()


def to_xlsx(dataset: dict[str, Any], metadata: dict[str, Any]) -> bytes:
    workbook = Workbook()
    data = workbook.active
    assert data is not None
    data.title = "Data"
    data.append(COLUMNS)
    for row in rows(dataset):
        data.append([row[column] for column in COLUMNS])
    fields = workbook.create_sheet("Fields")
    field_columns = ["id", "name", "section", "field_type", "unit", "per_arm", "required", "outcome", "timepoint"]
    fields.append(field_columns)
    for item in dataset["fields"]:
        fields.append(
            [str(item[column]) if isinstance(item[column], list | dict) else item[column] for column in field_columns]
        )
    studies = workbook.create_sheet("Studies")
    studies.append(["study_id", "label", "registry_ids", "reports", "arms"])
    for study in dataset["studies"]:
        studies.append(
            [
                study["study_id"],
                study["label"],
                ";".join(study["registry_ids"]),
                " | ".join(f"{r['title']} ({r['doi'] or 'no DOI'})" for r in study["reports"]),
                "; ".join(arm["label"] for arm in study["arms"]),
            ]
        )
    about = workbook.create_sheet("About")
    for key, value in metadata.items():
        about.append([key, str(value)])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


R_LOADER = """# Loads the OmniReview extraction dataset (long format: one row per value component).
dataset <- read.csv("dataset.csv", stringsAsFactors = FALSE)
dataset$value_numeric <- suppressWarnings(as.numeric(dataset$value))

# Example: arm-level continuous outcomes in wide format (mean, sd, n per study and arm).
continuous <- subset(dataset, field_type == "continuous")
if (nrow(continuous) > 0) {
  wide <- reshape(
    continuous[, c("study", "field", "arm", "component", "value_numeric")],
    idvar = c("study", "field", "arm"), timevar = "component", direction = "wide"
  )
  print(head(wide))
}
"""

PYTHON_LOADER = '''"""Loads the OmniReview extraction dataset (long format: one row per value component)."""

import pandas as pd

dataset = pd.read_csv("dataset.csv")
dataset["value_numeric"] = pd.to_numeric(dataset["value"], errors="coerce")

# Example: arm-level continuous outcomes in wide format (mean, sd, n per study and arm).
continuous = dataset[dataset["field_type"] == "continuous"]
if not continuous.empty:
    wide = continuous.pivot_table(
        index=["study", "field", "arm"], columns="component", values="value_numeric", aggfunc="first"
    )
    print(wide.head())
'''


def bundle(dataset: dict[str, Any], metadata: dict[str, Any]) -> bytes:
    readme = "\n".join(
        [
            "OmniReview extraction dataset",
            "",
            *(f"{key}: {value}" for key, value in metadata.items()),
            "",
            "dataset.csv    long format, one row per value component",
            "dataset.json   the full dataset with fields, studies, reports, arms, and provenance",
            "dataset.xlsx   the same data as spreadsheets",
            "load_dataset.R and load_dataset.py   loader scripts",
        ]
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README.txt", readme)
        archive.writestr("dataset.csv", to_csv(dataset))
        archive.writestr("dataset.json", to_json(dataset, metadata))
        archive.writestr("dataset.xlsx", to_xlsx(dataset, metadata))
        archive.writestr("load_dataset.R", R_LOADER)
        archive.writestr("load_dataset.py", PYTHON_LOADER)
    return buffer.getvalue()
