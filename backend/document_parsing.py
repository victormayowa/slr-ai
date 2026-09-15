"""Turn stored full texts into passages ("spans") that keep their kind, section, and page.

JATS XML (the format of PubMed Central and Europe PMC full texts) keeps the article's structure exactly. Born-digital
PDFs are read with pypdf one page at a time, with headings recognized from common section names; scanned PDFs have no
text layer and need OCR, which isn't available yet. DOCX and plain text files are also read. Each span's offsets refer
to the document's plain text, the span texts joined by SPAN_SEPARATOR, so quotes in AI output can be located later.
"""

import io
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Any, Literal

import pypdf
from defusedxml import DefusedXmlException
from defusedxml import ElementTree as SafeElementTree
from docx import Document as open_docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

PARSER_VERSION = "1"
SPAN_SEPARATOR = "\n\n"
MAX_SPAN_CHARS = 2000
MAX_PDF_PAGES = 2000
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
# With less text than this, a PDF is taken to be a scanned image without a text layer.
MIN_PDF_TEXT_CHARS = 200

FileKind = Literal["jats", "pdf", "docx", "text", "unsupported"]
MEDIA_TYPES: dict[str, str] = {
    "jats": "application/xml",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text": "text/plain",
    "unsupported": "application/octet-stream",
}
READABLE_FORMATS = "JATS XML, PDF with a text layer, DOCX, and plain text"


class ParseError(Exception):
    """A file can't be read. The message is safe to show users."""


@dataclass
class ParsedSpan:
    # "title", "abstract", "heading", "paragraph", "table", "caption", or "reference"
    kind: str
    text: str
    section: str = ""
    page: int | None = None
    # A table or figure label, such as "Table 2", or a reference number.
    label: str = ""


@dataclass
class ParsedDocument:
    parser: str
    spans: list[ParsedSpan]
    page_count: int | None = None


def span_offsets(texts: Iterable[str]) -> list[tuple[int, int]]:
    """Start and end of each text in the texts joined by SPAN_SEPARATOR."""
    offsets, position = [], 0
    for text in texts:
        offsets.append((position, position + len(text)))
        position += len(text) + len(SPAN_SEPARATOR)
    return offsets


def detect_kind(file_name: str, content: bytes) -> FileKind:
    name = file_name.lower()
    head = content[:4096].lstrip(b"\xef\xbb\xbf \t\r\n")
    if b"%PDF-" in content[:1024]:
        return "pdf"
    if content.startswith(b"PK\x03\x04") and name.endswith(".docx"):
        return "docx"
    if head.startswith(b"<") and b"<html" not in head.lower() and re.search(rb"<article[\s>]", content[:20_000]):
        return "jats"
    if name.endswith((".txt", ".text")):
        return "text"
    return "unsupported"


def parse_document(content: bytes, kind: FileKind) -> ParsedDocument:
    if kind == "jats":
        return _parse_jats(content)
    if kind == "pdf":
        return _parse_pdf(content)
    if kind == "docx":
        return _parse_docx(content)
    if kind == "text":
        return _parse_text(content)
    raise ParseError(f"This file type is stored but not read. Readable formats: {READABLE_FORMATS}.")


def _chunks(text: str) -> list[str]:
    """Split long text at sentence ends into pieces of at most MAX_SPAN_CHARS."""
    if len(text) <= MAX_SPAN_CHARS:
        return [text] if text else []
    chunks: list[str] = []
    current = ""
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        while len(sentence) > MAX_SPAN_CHARS:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(sentence[:MAX_SPAN_CHARS])
            sentence = sentence[MAX_SPAN_CHARS:]
        if current and len(current) + 1 + len(sentence) > MAX_SPAN_CHARS:
            chunks.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}" if current else sentence
    if current:
        chunks.append(current)
    return chunks


# --- JATS XML ---

# Elements whose text runs on within a paragraph. Other elements are separated by a space.
_INLINE = frozenset(
    {
        "italic",
        "bold",
        "sup",
        "sub",
        "sc",
        "underline",
        "monospace",
        "roman",
        "sans-serif",
        "strike",
        "overline",
        "xref",
        "ext-link",
        "uri",
        "email",
        "named-content",
        "styled-content",
        "inline-formula",
        "abbrev",
        "math",
    }
)
# Blocks that can sit inside a paragraph; they become spans of their own.
_BLOCKS_IN_PARAGRAPH = frozenset(
    {"table-wrap", "fig", "boxed-text", "list", "disp-quote", "supplementary-material", "table-wrap-group", "fig-group"}
)
_SKIPPED_ABSTRACTS = ("graphical", "teaser", "toc")
# In a mixed-citation the punctuation between these parts is written out, so they need no extra spaces.
_MIXED_CITATION_INLINE = _INLINE | frozenset(
    {
        "person-group",
        "name",
        "string-name",
        "surname",
        "given-names",
        "prefix",
        "suffix",
        "etal",
        "collab",
        "article-title",
        "chapter-title",
        "source",
        "year",
        "month",
        "day",
        "volume",
        "issue",
        "fpage",
        "lpage",
        "page-range",
        "elocation-id",
        "edition",
        "publisher-name",
        "publisher-loc",
        "pub-id",
        "comment",
        "date-in-citation",
    }
)


def _local(tag: Any) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _first(element: Any, name: str) -> Any:
    return next((child for child in element if _local(child.tag) == name), None)


def _text(element: Any, skip: frozenset[str] = frozenset(), inline: frozenset[str] = _INLINE) -> str:
    if element is None:
        return ""
    parts: list[str] = []

    def walk(node: Any) -> None:
        if node.text:
            parts.append(node.text)
        for child in node:
            tag = _local(child.tag)
            if tag not in skip:
                spacer = "" if tag in inline else " "
                parts.append(spacer)
                walk(child)
                parts.append(spacer)
            if child.tail:
                parts.append(child.tail)

    walk(element)
    return " ".join("".join(parts).split())


class _JatsWalker:
    def __init__(self) -> None:
        self.spans: list[ParsedSpan] = []

    def add(self, kind: str, text: str, section: str, label: str = "") -> None:
        for chunk in _chunks(text):
            self.spans.append(ParsedSpan(kind, chunk, section, None, label))

    def blocks(self, element: Any, section: str, paragraph_kind: str = "paragraph") -> None:
        for child in element:
            self.block(child, section, paragraph_kind)

    def block(self, child: Any, section: str, paragraph_kind: str) -> None:
        tag = _local(child.tag)
        if tag in ("sec", "ack", "app"):
            heading = _text(_first(child, "title")) or ("Acknowledgements" if tag == "ack" else "")
            inner = f"{section} > {heading}" if section and heading else (heading or section)
            if heading:
                self.add("heading", heading, inner)
            self.blocks(child, inner, paragraph_kind)
        elif tag == "p":
            self.add(paragraph_kind, _text(child, _BLOCKS_IN_PARAGRAPH), section)
            self.nested_blocks(child, section, paragraph_kind)
        elif tag == "list":
            for item in child:
                if _local(item.tag) == "list-item":
                    self.add(paragraph_kind, _text(item, _BLOCKS_IN_PARAGRAPH), section)
                    self.nested_blocks(item, section, paragraph_kind)
        elif tag == "table-wrap":
            self.table(child, section)
        elif tag in ("fig", "supplementary-material"):
            self.add("caption", _text(_first(child, "caption")), section, _text(_first(child, "label")))
        elif tag in ("boxed-text", "disp-quote", "statement", "fig-group", "table-wrap-group", "app-group"):
            self.blocks(child, section, paragraph_kind)
        elif tag == "ref-list":
            self.references(child)

    def nested_blocks(self, element: Any, section: str, paragraph_kind: str) -> None:
        for child in element:
            if _local(child.tag) in _BLOCKS_IN_PARAGRAPH:
                self.block(child, section, paragraph_kind)
            else:
                self.nested_blocks(child, section, paragraph_kind)

    def table(self, wrap: Any, section: str) -> None:
        label = _text(_first(wrap, "label"))
        self.add("caption", _text(_first(wrap, "caption")), section, label)
        rows = []
        for row in wrap.iter():
            if _local(row.tag) == "tr":
                cells = [_text(cell) for cell in row if _local(cell.tag) in ("td", "th")]
                if any(cells):
                    rows.append(" | ".join(cells))
        if footer := _text(_first(wrap, "table-wrap-foot")):
            rows.append(footer)
        if rows:
            self.spans.append(ParsedSpan("table", "\n".join(rows), section, None, label))

    def references(self, ref_list: Any) -> None:
        heading = _text(_first(ref_list, "title")) or "References"
        self.add("heading", heading, heading)
        for child in ref_list:
            tag = _local(child.tag)
            if tag == "ref":
                mixed = _first(child, "mixed-citation")
                if mixed is not None:
                    text = _text(mixed, inline=_MIXED_CITATION_INLINE)
                else:
                    text = _text(child, frozenset({"label"}))
                if text:
                    self.spans.append(ParsedSpan("reference", text, heading, None, _text(_first(child, "label"))))
            elif tag == "ref-list":
                self.references(child)


def _parse_jats(content: bytes) -> ParsedDocument:
    try:
        root = SafeElementTree.fromstring(content)
    except (SafeElementTree.ParseError, DefusedXmlException) as exc:
        raise ParseError("This XML file couldn't be read") from exc
    article = (
        root
        if _local(root.tag) == "article"
        else next((element for element in root.iter() if _local(element.tag) == "article"), None)
    )
    if article is None:
        raise ParseError("This XML file isn't a JATS article")

    walker = _JatsWalker()
    front = _first(article, "front")
    meta = _first(front, "article-meta") if front is not None else None
    if meta is not None:
        title_group = _first(meta, "title-group")
        if title_group is not None:
            walker.add("title", _text(_first(title_group, "article-title")), "")
        for abstract in meta:
            if _local(abstract.tag) == "abstract" and abstract.get("abstract-type") not in _SKIPPED_ABSTRACTS:
                walker.blocks(abstract, _text(_first(abstract, "title")) or "Abstract", "abstract")
    for part in ("body", "back", "floats-group"):
        element = _first(article, part)
        if element is not None:
            walker.blocks(element, "")

    if not any(span.kind in ("paragraph", "abstract", "table") for span in walker.spans):
        raise ParseError("No article text was found in this XML file")
    return ParsedDocument(f"jats/v{PARSER_VERSION}", walker.spans)


# --- PDF, DOCX, and plain text ---

_SECTION_NAMES = (
    r"abstract|summary|background|introduction|methods?|materials and methods|patients and methods|study design|"
    r"results|findings|discussion|conclusions?|interpretation|limitations|strengths and limitations|references|"
    r"bibliography|acknowledge?ments?|funding|conflicts? of interest|competing interests|declarations?|"
    r"supplementary (?:material|materials|data|information)|appendix|appendices"
)
_HEADING = re.compile(rf"(?:\d+(?:\.\d+)*\.?\s+)?({_SECTION_NAMES})\s*:?", re.IGNORECASE)
_REFERENCE_SECTION = re.compile(r"references|bibliography", re.IGNORECASE)


def _heading(line: str) -> str | None:
    if len(line) > 60:
        return None
    match = _HEADING.fullmatch(line.strip())
    return match.group(1).strip().capitalize() if match else None


def _paragraph_kind(section: str) -> str:
    return "reference" if _REFERENCE_SECTION.fullmatch(section) else "paragraph"


def _line_spans(text: str, page: int | None, section: str) -> tuple[list[ParsedSpan], str]:
    """Spans from extracted lines: blank lines end paragraphs, and lines naming a common section are headings."""
    text = re.sub(r"(\w)-\n([a-z])", r"\1\2", text)
    spans: list[ParsedSpan] = []
    lines: list[str] = []

    def flush() -> None:
        paragraph = " ".join(" ".join(lines).split())
        spans.extend(ParsedSpan(_paragraph_kind(section), chunk, section, page) for chunk in _chunks(paragraph))
        lines.clear()

    for raw in text.splitlines():
        line = raw.strip()
        heading = _heading(line) if line else None
        if not line:
            flush()
        elif heading:
            flush()
            section = heading
            spans.append(ParsedSpan("heading", heading, section, page))
        else:
            lines.append(line)
    flush()
    return spans, section


def _parse_pdf(content: bytes) -> ParsedDocument:
    spans: list[ParsedSpan] = []
    section = ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParseError("This PDF is password protected")
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise ParseError(f"This PDF has more than {MAX_PDF_PAGES} pages")
        for number, page in enumerate(reader.pages, start=1):
            page_spans, section = _line_spans(page.extract_text() or "", number, section)
            spans.extend(page_spans)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("This PDF couldn't be read; it may be damaged or password protected") from exc
    if sum(len(span.text) for span in spans) < MIN_PDF_TEXT_CHARS:
        raise ParseError(
            "This PDF has almost no text layer, so it is probably a scanned image. OCR isn't available yet: "
            "upload a text-based PDF or the XML version if there is one."
        )
    return ParsedDocument(f"pypdf-{pypdf.__version__}/v{PARSER_VERSION}", spans, page_count)


def _parse_docx(content: bytes) -> ParsedDocument:
    spans: list[ParsedSpan] = []
    section = ""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ParseError("This DOCX file is too large once uncompressed")
        document = open_docx(io.BytesIO(content))
        for child in document.element.body.iterchildren():
            tag = _local(child.tag)
            if tag == "p":
                paragraph = DocxParagraph(child, document)
                text = " ".join(paragraph.text.split())
                if not text:
                    continue
                style = (paragraph.style.name if paragraph.style is not None else "") or ""
                if style == "Title":
                    spans.append(ParsedSpan("title", text))
                elif style.startswith("Heading") or _heading(text):
                    section = text
                    spans.append(ParsedSpan("heading", text, section))
                else:
                    spans.extend(ParsedSpan(_paragraph_kind(section), chunk, section) for chunk in _chunks(text))
            elif tag == "tbl":
                table = DocxTable(child, document)
                rows = [" | ".join(" ".join(cell.text.split()) for cell in row.cells) for row in table.rows]
                if any(row.strip(" |") for row in rows):
                    spans.append(ParsedSpan("table", "\n".join(rows), section))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("This DOCX file couldn't be read") from exc
    if not spans:
        raise ParseError("No text was found in this DOCX file")
    return ParsedDocument(f"docx/v{PARSER_VERSION}", spans)


def _parse_text(content: bytes) -> ParsedDocument:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    spans, _ = _line_spans(text, None, "")
    if not spans:
        raise ParseError("No text was found in this file")
    return ParsedDocument(f"text/v{PARSER_VERSION}", spans)
