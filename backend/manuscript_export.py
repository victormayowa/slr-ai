"""Manuscript export with Pandoc: Word, LaTeX (with figures, as a zip), PDF (Tectonic), and Markdown (zip).

Evidence markers are removed, citations are rendered in the chosen style (built-in styles by OmniReview, other CSL
styles by Pandoc's citeproc with the style fetched from the CSL styles repository), and embedded tables and figures are
built from the review's current records. Without Pandoc, Word export falls back to a simpler document.
"""

import asyncio
import io
import json
import os
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import requests
from docx import Document as WordDocument
from docx.shared import Inches
from sqlalchemy import select
from sqlalchemy.orm import Session

import citations
import models
from claims import CITATION, CITATION_ID, EMBED, EVIDENCE
from manuscript_assets import Table, figure, table
from publishing_tools import tool_path

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CSL_REPOSITORY = "https://raw.githubusercontent.com/citation-style-language/styles/master"
CSL_CACHE = Path.home() / ".cache" / "omnireview" / "csl"
PANDOC_TIMEOUT_SECONDS = int(os.getenv("PANDOC_TIMEOUT_SECONDS", "600"))


class ExportUnavailable(Exception):
    """The export can't be produced (a tool is missing or failed). The message is safe to show users."""


@dataclass
class BuiltDocument:
    markdown: str
    files: dict[str, bytes] = field(default_factory=dict)
    csl_references: list[dict[str, Any]] = field(default_factory=list)
    style: str = "vancouver"


def fetch_csl_style(style_id: str) -> Path:
    if not citations.CSL_STYLE_ID.match(style_id):
        raise ExportUnavailable("Citation style ids use lowercase letters, digits, and hyphens")
    path = CSL_CACHE / f"{style_id}.csl"
    if path.exists():
        return path
    try:
        response = requests.get(f"{CSL_REPOSITORY}/{style_id}.csl", timeout=30)
    except requests.RequestException as exc:
        raise ExportUnavailable("The citation style couldn't be downloaded. Try again shortly.") from exc
    if response.status_code != 200 or "<style" not in response.text:
        raise ExportUnavailable(f"The CSL styles repository has no style named {style_id}")
    CSL_CACHE.mkdir(parents=True, exist_ok=True)
    path.write_text(response.text)
    return path


def clean_line(line: str) -> str:
    return re.sub(r" {2,}", " ", re.sub(r"\s+([.,;:])", r"\1", EVIDENCE.sub("", line))).rstrip()


def markdown_table(item: Table, number: int) -> str:
    def cell(value: object) -> str:
        return str(value).replace("|", "\\|").replace("\n", " ")

    lines = [f"**Table {number}. {item.title}**", "", "| " + " | ".join(cell(h) for h in item.header) + " |"]
    lines.append("|" + "---|" * len(item.header))
    rows = item.rows or [["" for _ in item.header]]
    lines += ["| " + " | ".join(cell(c) for c in row) + " |" for row in rows]
    if item.note:
        lines += ["", item.note]
    return "\n".join(lines)


def references_for(db: Session, project_id: int) -> dict[int, dict[str, Any]]:
    rows = db.scalars(
        select(models.ManuscriptReference)
        .where(models.ManuscriptReference.project_id == project_id)
        .order_by(models.ManuscriptReference.id)
    )
    return {row.id: row.csl for row in rows}


def _yaml(value: object) -> str:
    return json.dumps(value, ensure_ascii=False)


async def build_document(
    db: Session, project: models.Project, manuscript: models.Manuscript, fmt: str
) -> BuiltDocument:
    references = references_for(db, project.id)
    style = manuscript.citation_style
    built_in = style in citations.BUILT_IN_STYLES
    sections = sorted(manuscript.sections, key=lambda s: s.position)
    texts = [s.content for s in sections]
    bibliography: list[tuple[int | None, int, str]] = []
    if built_in:
        rendered = citations.render(texts, references, style)
        texts, bibliography = rendered.texts, rendered.bibliography
    else:
        texts = [
            CITATION.sub(lambda m: "[" + "; ".join(f"@ref{i}" for i in CITATION_ID.findall(m.group(0))) + "]", text)
            for text in texts
        ]
    preferred = {"docx": ("png", "svg"), "tex": ("pdf", "png", "svg"), "pdf": ("pdf", "png", "svg")}.get(
        fmt, ("svg", "png")
    )
    files: dict[str, bytes] = {}
    tables = figures = 0
    lines = [
        "---",
        f"title: {_yaml(manuscript.title)}",
        "author:",
        *[f"  - {_yaml(author.name)}" for author in manuscript.authors],
    ]
    if manuscript.keywords:
        lines.append(f"keywords: {_yaml(manuscript.keywords)}")
    lines += ["---", ""]
    affiliations = [a for a in manuscript.authors if a.affiliation]
    if affiliations:
        lines += [
            f"{a.name}: {a.affiliation}" + (f" (ORCID {a.orcid})" if a.orcid else "") + "  " for a in affiliations
        ]
        lines.append("")
    corresponding = next((a for a in manuscript.authors if a.corresponding), None)
    if corresponding:
        lines += [
            f"Corresponding author: {corresponding.name}" + (f", {corresponding.email}" if corresponding.email else ""),
            "",
        ]
    if manuscript.keywords:
        lines += [f"Keywords: {', '.join(manuscript.keywords)}", ""]
    for section, text in zip(sections, texts, strict=True):
        if not text.strip():
            continue
        lines += [f"# {section.title}", ""]
        for line in text.splitlines():
            embed = EMBED.match(line.strip())
            if embed is None:
                lines.append(clean_line(line))
                continue
            kind, key = embed.groups()
            if kind == "table":
                item = table(db, project, key)
                if item is not None:
                    tables += 1
                    lines += ["", markdown_table(item, tables), ""]
                continue
            picture = await figure(db, project, key)
            best = picture.best(preferred) if picture else None
            if picture is not None and best is not None:
                figures += 1
                name = f"figure-{figures}.{best[0]}"
                files[name] = best[1]
                lines += ["", f"![Figure {figures}. {picture.title}]({name})", ""]
        lines.append("")
    if built_in and bibliography:
        lines += ["# References", ""]
        for number, _, entry in bibliography:
            lines.append(f"{number}. {entry}" if number is not None else f"{entry}\n")
    elif not built_in and references:
        lines += ["# References", "", "::: {#refs}", ":::"]
    csl_references = [{**csl, "id": f"ref{ref_id}"} for ref_id, csl in references.items()]
    return BuiltDocument("\n".join(lines) + "\n", files, csl_references, style)


def _zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def _simple_docx(document: BuiltDocument) -> bytes:
    """A plain Word document when Pandoc isn't installed: headings, paragraphs, tables, and PNG figures."""
    word = WordDocument()
    lines = document.markdown.split("\n---\n", 1)[-1].splitlines()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        rows = [r for r in table_rows if not all(set(c) <= {"-"} for c in r)]
        grid = word.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        grid.style = "Table Grid"
        for row, values in zip(grid.rows, rows, strict=False):
            for cell, value in zip(row.cells, values, strict=False):
                cell.text = value
        table_rows.clear()

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            table_rows.append([c.strip().replace("\\|", "|") for c in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped or stripped.startswith(":::"):
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            word.add_heading(stripped.lstrip("#").strip(), level=min(level, 4))
        elif (image := re.match(r"!\[(.*)\]\((.*)\)", stripped)) is not None:
            name = image.group(2)
            if name.endswith(".png") and name in document.files:
                word.add_picture(io.BytesIO(document.files[name]), width=Inches(6))
            word.add_paragraph(image.group(1))
        else:
            word.add_paragraph(stripped.replace("**", "").replace("*", ""))
    flush_table()
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _run_pandoc(document: BuiltDocument, fmt: str) -> bytes:
    pandoc = tool_path("PANDOC_PATH", "pandoc")
    if pandoc is None:
        if fmt == "docx":
            return _simple_docx(document)
        raise ExportUnavailable("Pandoc isn't installed on the server. Run backend/scripts/setup_publishing_tools.sh.")
    with tempfile.TemporaryDirectory(prefix="omnireview-export-") as folder:
        root = Path(folder)
        (root / "manuscript.md").write_text(document.markdown)
        for name, content in document.files.items():
            (root / name).write_bytes(content)
        args = [pandoc, "manuscript.md", "--from", "markdown", "--resource-path", folder]
        if document.style not in citations.BUILT_IN_STYLES and document.csl_references:
            (root / "references.json").write_text(json.dumps(document.csl_references))
            args += ["--citeproc", "--bibliography", "references.json", "--csl", str(fetch_csl_style(document.style))]
        env = dict(os.environ)
        tool_dirs = {str(Path(p).parent) for p in (pandoc, tool_path("RSVG_CONVERT_PATH", "rsvg-convert")) if p}
        env["PATH"] = os.pathsep.join([*tool_dirs, env.get("PATH", "")])
        if fmt == "docx":
            output = "manuscript.docx"
        elif fmt == "tex":
            output = "manuscript.tex"
            args.append("--standalone")
        else:
            tectonic = tool_path("TECTONIC_PATH", "tectonic")
            if tectonic is None:
                raise ExportUnavailable("PDF export needs Tectonic. Run backend/scripts/setup_publishing_tools.sh.")
            output = "manuscript.pdf"
            args += ["--pdf-engine", tectonic, "-V", "geometry:margin=2.5cm"]
        args += ["--output", output]
        try:
            completed = subprocess.run(
                args, cwd=folder, capture_output=True, text=True, timeout=PANDOC_TIMEOUT_SECONDS, env=env, check=False
            )
        except subprocess.TimeoutExpired as exc:
            raise ExportUnavailable("The export took too long and was stopped") from exc
        if completed.returncode != 0 or not (root / output).exists():
            raise ExportUnavailable(f"Pandoc couldn't build the {fmt.upper()} file: {completed.stderr.strip()[-500:]}")
        content = (root / output).read_bytes()
        if fmt == "tex":
            return _zip({"manuscript.tex": content, **document.files})
        return content


async def export_manuscript(
    db: Session, project: models.Project, manuscript: models.Manuscript, fmt: str
) -> tuple[bytes, str, str]:
    document = await build_document(db, project, manuscript, fmt)
    if fmt == "md":
        return (
            _zip({"manuscript.md": document.markdown.encode(), **document.files}),
            "application/zip",
            "manuscript-markdown.zip",
        )
    content = await asyncio.to_thread(_run_pandoc, document, fmt)
    return {
        "docx": (content, DOCX_TYPE, "manuscript.docx"),
        "tex": (content, "application/zip", "manuscript-latex.zip"),
        "pdf": (content, "application/pdf", "manuscript.pdf"),
    }[fmt]
