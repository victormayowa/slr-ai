"""Certainty of evidence and interpretation: GRADE assessments per outcome with suggested ratings from the analyses,
summary of findings tables with absolute effects, Evidence to Decision frameworks, comparison with prior reviews
(study overlap and changed conclusions), and interpretive text (informative statements, limitations, and an AI
plain-language summary whose numbers are checked against the summary of findings) approved by a clinical expert.
"""

import csv
import io
import re
from typing import Any, Literal

from docx import Document
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

import grading
import models
from ai_access import new_ai_run, project_ai, record_usage
from audit import record_event
from citation_chasing import resolve_openalex_ids
from database import get_db
from extraction_data import included_studies
from llm.prompts import PLAIN_LANGUAGE_PROMPT
from llm.runner import complete_structured
from permissions import Permission
from projects_routes import ProjectAccess, get_in_project, project_access
from rate_limiting import ai_rate_limit
from services.errors import LLMError, SearchError
from services.openalex import openalex_raw_works
from synthesis_routes import final_run
from workflow import require_stage_open

router = APIRouter(prefix="/api/projects/{project_id}", tags=["certainty"])

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEASURE_LABELS = {"RR": "RR", "OR": "OR", "HR": "HR", "RD": "RD", "MD": "MD", "SMD": "SMD", "ROM": "Ratio of means"}
DIRECTIONS = ("favours_intervention", "favours_comparator", "no_difference", "uncertain")


def _audit(db: Session, access: ProjectAccess, action: str, entity_type: str, entity_id: int, details: dict) -> None:
    record_event(
        db,
        project_id=access.project.id,
        actor_id=access.user.id,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        details=details,
    )


def intervention_name(project: models.Project) -> str:
    elements = (project.protocol.question_elements if project.protocol else None) or {}
    return (
        elements.get("intervention") or elements.get("exposure") or elements.get("index_test") or "the intervention"
    ).strip()


def _fmt(value: float | None, digits: int = 2) -> str:
    return "–" if value is None else f"{value:.{digits}f}"


def effect_summary(run: models.AnalysisRun) -> dict[str, Any]:
    """The main estimate of a run, whatever the analysis type, on the analysis scale and (for ratios) exponentiated."""
    results = run.results or {}
    summary = dict(results.get("summary") or results.get("two_stage") or {})
    if "mu_median" in summary:
        summary.update(
            estimate=summary.get("mu_median"),
            ci_lower=summary.get("mu_lower"),
            ci_upper=summary.get("mu_upper"),
            exp_estimate=summary.get("exp_mu_median"),
            exp_ci_lower=summary.get("exp_mu_lower"),
            exp_ci_upper=summary.get("exp_mu_upper"),
        )
    rows = run.dataset.get("rows", [])
    summary["studies"] = len({row.get("study_id") for row in rows})
    summary["participants"] = (
        int(sum((row.get("n1i") or 0) + (row.get("n2i") or 0) + (row.get("n") or 0) for row in rows)) or None
    )
    return summary


def analysis_for(
    db: Session, assessment: models.GradeAssessment
) -> tuple[models.Analysis | None, models.AnalysisRun | None]:
    analysis = db.get(models.Analysis, assessment.analysis_id) if assessment.analysis_id else None
    if analysis is None:
        return None, None
    return analysis, final_run(analysis) or next((r for r in reversed(analysis.runs) if r.status == "succeeded"), None)


def sof_row(db: Session, project: models.Project, assessment: models.GradeAssessment) -> dict[str, Any]:
    analysis, run = analysis_for(db, assessment)
    row: dict[str, Any] = {
        "grade_assessment_id": assessment.id,
        "outcome": assessment.outcome,
        "comparison": assessment.comparison,
        "importance": assessment.importance,
        "certainty": assessment.certainty,
        "certainty_label": grading.LEVEL_LABELS[assessment.certainty],
        "status": assessment.status,
        "analysis_id": analysis.id if analysis else None,
        "run_id": run.id if run else None,
        "final": bool(run and run.is_final),
        "studies": None,
        "participants": None,
        "relative_effect": "",
        "absolute_effects": [],
        "footnotes": [
            f"{grading.DOWNGRADE_DOMAINS.get(key) or grading.UPGRADE_DOMAINS.get(key, key)}: {entry.get('rationale')}"
            for key, entry in (assessment.domains or {}).items()
            if int(entry.get("rating", 0)) != 0 and entry.get("rationale")
        ],
    }
    direction = "unknown"
    if run is not None:
        measure = run.spec.get("measure", "")
        summary = effect_summary(run)
        row["studies"], row["participants"] = summary["studies"], summary["participants"]
        ratio = measure in grading.RATIO
        estimate = summary.get("exp_estimate") if ratio else summary.get("estimate")
        lower = summary.get("exp_ci_lower") if ratio else summary.get("ci_lower")
        upper = summary.get("exp_ci_upper") if ratio else summary.get("ci_upper")
        if estimate is not None:
            row["relative_effect"] = (
                f"{MEASURE_LABELS.get(measure, measure)} {_fmt(estimate)} ({_fmt(lower)} to {_fmt(upper)})"
            )
        first_absolute = None
        if estimate is not None and lower is not None and upper is not None:
            for baseline in assessment.baseline_risks or []:
                absolute = grading.absolute_effects(measure, estimate, lower, upper, float(baseline.get("risk", 0)))
                if absolute:
                    row["absolute_effects"].append({"label": baseline.get("label", ""), **absolute})
                    first_absolute = first_absolute or absolute
        direction = grading.effect_direction(measure, summary, first_absolute, assessment.mid, assessment.mid_scale)
    row["direction"] = direction
    row["informative_statement"] = grading.informative_statement(
        assessment.certainty, direction, intervention_name(project), assessment.outcome.lower()
    )
    return row


def grade_out(db: Session, project: models.Project, assessment: models.GradeAssessment) -> dict[str, Any]:
    analysis, run = analysis_for(db, assessment)
    suggestions = {}
    if run is not None and analysis is not None:
        suggestions = grading.suggested_ratings(
            analysis.analysis_type,
            run.spec.get("measure", ""),
            run.results or {},
            run.dataset.get("rows", []),
            assessment.mid,
            assessment.mid_scale,
            assessment.starting_certainty,
        )
    return {
        "id": assessment.id,
        "outcome": assessment.outcome,
        "comparison": assessment.comparison,
        "analysis_id": assessment.analysis_id,
        "importance": assessment.importance,
        "starting_certainty": assessment.starting_certainty,
        "domains": assessment.domains,
        "certainty": assessment.certainty,
        "mid": assessment.mid,
        "mid_scale": assessment.mid_scale,
        "outcome_direction": assessment.outcome_direction,
        "baseline_risks": assessment.baseline_risks,
        "status": assessment.status,
        "signed_off_by": assessment.signed_off_by.full_name if assessment.signed_off_by else None,
        "signed_off_at": assessment.signed_off_at,
        "sign_off_note": assessment.sign_off_note,
        "suggestions": suggestions,
        "summary_of_findings": sof_row(db, project, assessment),
    }


def _assessments(db: Session, project_id: int) -> list[models.GradeAssessment]:
    return list(
        db.scalars(
            select(models.GradeAssessment)
            .where(models.GradeAssessment.project_id == project_id)
            .order_by(models.GradeAssessment.id)
        )
    )


@router.get("/certainty/catalog")
def certainty_catalog(access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT))):
    return {
        "levels": [{"key": k, "label": grading.LEVEL_LABELS[k]} for k in grading.LEVELS],
        "downgrade_domains": grading.DOWNGRADE_DOMAINS,
        "upgrade_domains": grading.UPGRADE_DOMAINS,
        "etd_criteria": grading.ETD_CRITERIA,
        "recommendation_types": grading.RECOMMENDATION_TYPES,
        "conclusion_fields": grading.CONCLUSION_FIELDS,
        "conclusion_directions": DIRECTIONS,
    }


@router.get("/certainty/outcomes")
def certainty_outcomes(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    """Outcomes with an approved analysis and final results, and the GRADE assessment of each."""
    graded = {(a.outcome.casefold(), a.analysis_id): a.id for a in _assessments(db, access.project.id)}
    analyses = db.scalars(
        select(models.Analysis).where(models.Analysis.project_id == access.project.id).order_by(models.Analysis.id)
    )
    return [
        {
            "analysis_id": analysis.id,
            "title": analysis.title,
            "outcome": analysis.outcome,
            "analysis_type": analysis.analysis_type,
            "final_run_id": final.id,
            "grade_assessment_id": graded.get((analysis.outcome.casefold(), analysis.id)),
        }
        for analysis in analyses
        if analysis.status == "approved" and (final := final_run(analysis)) is not None
    ]


class DomainRating(BaseModel):
    rating: int = Field(ge=-2, le=2)
    rationale: str = Field("", max_length=5000)


class BaselineRisk(BaseModel):
    label: str = Field(min_length=1, max_length=200)
    risk: float = Field(ge=0, le=1)


class GradeIn(BaseModel):
    outcome: str = Field(min_length=1, max_length=300)
    comparison: str = Field("", max_length=300)
    analysis_id: int | None = None
    importance: Literal["critical", "important", "not_important"] = "critical"
    starting_certainty: Literal["high", "low"] = "high"
    domains: dict[str, DomainRating] = Field(default_factory=dict)
    mid: float | None = Field(None, gt=0)
    mid_scale: Literal["", "per_1000", "units"] = ""
    outcome_direction: Literal["lower_is_better", "higher_is_better"] = "lower_is_better"
    baseline_risks: list[BaselineRisk] = Field(default_factory=list, max_length=5)


def _apply_grade(db: Session, access: ProjectAccess, assessment: models.GradeAssessment, body: GradeIn) -> None:
    known = set(grading.DOWNGRADE_DOMAINS) | set(grading.UPGRADE_DOMAINS)
    for key, entry in body.domains.items():
        if key not in known:
            raise HTTPException(status_code=422, detail=f"Unknown GRADE domain: {key}")
        if key in grading.DOWNGRADE_DOMAINS and entry.rating > 0:
            raise HTTPException(
                status_code=422, detail=f"{grading.DOWNGRADE_DOMAINS[key]} can only rate certainty down"
            )
        if key in grading.UPGRADE_DOMAINS and entry.rating < 0:
            raise HTTPException(status_code=422, detail=f"{grading.UPGRADE_DOMAINS[key]} can only rate certainty up")
        if entry.rating != 0 and len(entry.rationale.strip()) < 10:
            raise HTTPException(status_code=422, detail=f"Explain the rating for {key.replace('_', ' ')}")
    if body.analysis_id is not None:
        get_in_project(db, models.Analysis, body.analysis_id, access.project.id, "Analysis")
    if body.starting_certainty == "high" and any(
        body.domains.get(k, DomainRating(rating=0, rationale="")).rating for k in grading.UPGRADE_DOMAINS
    ):
        raise HTTPException(
            status_code=422, detail="Rating up applies to evidence starting at low certainty (non-randomized studies)"
        )
    assessment.outcome, assessment.comparison, assessment.analysis_id = (
        body.outcome.strip(),
        body.comparison.strip(),
        body.analysis_id,
    )
    assessment.importance, assessment.starting_certainty = body.importance, body.starting_certainty
    assessment.domains = {k: v.model_dump() for k, v in body.domains.items()}
    assessment.certainty = grading.certainty(body.starting_certainty, assessment.domains)
    assessment.mid, assessment.mid_scale, assessment.outcome_direction = (
        body.mid,
        body.mid_scale,
        body.outcome_direction,
    )
    assessment.baseline_risks = [b.model_dump() for b in body.baseline_risks]
    assessment.status, assessment.signed_off_by_id, assessment.signed_off_at, assessment.sign_off_note = (
        "draft",
        None,
        None,
        "",
    )


@router.get("/grade")
def list_grade(access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)):
    return [grade_out(db, access.project, a) for a in _assessments(db, access.project.id)]


@router.post("/grade", status_code=201)
def create_grade(
    body: GradeIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    duplicate = db.scalar(
        select(models.GradeAssessment.id).where(
            models.GradeAssessment.project_id == access.project.id,
            models.GradeAssessment.outcome == body.outcome.strip(),
            models.GradeAssessment.comparison == body.comparison.strip(),
        )
    )
    if duplicate is not None:
        raise HTTPException(status_code=409, detail="This outcome and comparison already has a GRADE assessment")
    assessment = models.GradeAssessment(
        project_id=access.project.id, outcome=body.outcome, created_by_id=access.user.id
    )
    _apply_grade(db, access, assessment, body)
    db.add(assessment)
    db.flush()
    _audit(
        db,
        access,
        "grade.created",
        "grade_assessment",
        assessment.id,
        {"outcome": assessment.outcome, "certainty": assessment.certainty},
    )
    db.commit()
    return grade_out(db, access.project, assessment)


@router.get("/grade/{assessment_id}")
def get_grade(
    assessment_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    return grade_out(
        db,
        access.project,
        get_in_project(db, models.GradeAssessment, assessment_id, access.project.id, "GRADE assessment"),
    )


@router.put("/grade/{assessment_id}")
def update_grade(
    assessment_id: int,
    body: GradeIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    """Change ratings or inputs. A signed-off assessment returns to draft."""
    require_stage_open(db, access.project.id, "certainty")
    assessment = get_in_project(db, models.GradeAssessment, assessment_id, access.project.id, "GRADE assessment")
    before = {"domains": assessment.domains, "certainty": assessment.certainty, "status": assessment.status}
    _apply_grade(db, access, assessment, body)
    _audit(
        db,
        access,
        "grade.updated",
        "grade_assessment",
        assessment.id,
        {"before": before, "domains": assessment.domains, "certainty": assessment.certainty},
    )
    db.commit()
    return grade_out(db, access.project, assessment)


class SignOffIn(BaseModel):
    note: str = Field(min_length=10, max_length=5000)


@router.post("/grade/{assessment_id}/sign-off")
def sign_off_grade(
    assessment_id: int,
    body: SignOffIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    assessment = get_in_project(db, models.GradeAssessment, assessment_id, access.project.id, "GRADE assessment")
    unrated = [label for key, label in grading.DOWNGRADE_DOMAINS.items() if key not in (assessment.domains or {})]
    if unrated:
        raise HTTPException(
            status_code=409,
            detail=f"Rate every domain before signing off (even when not rated down): {', '.join(unrated)}",
        )
    analysis, run = analysis_for(db, assessment)
    if analysis is not None and (run is None or not run.is_final):
        raise HTTPException(
            status_code=409,
            detail="The linked analysis has no final results (approved analysis run on the locked data set)",
        )
    assessment.status, assessment.signed_off_by_id, assessment.signed_off_at = (
        "signed_off",
        access.user.id,
        models.utcnow(),
    )
    assessment.sign_off_note = body.note.strip()
    _audit(
        db,
        access,
        "grade.signed_off",
        "grade_assessment",
        assessment.id,
        {
            "certainty": assessment.certainty,
            "domains": assessment.domains,
            "note": assessment.sign_off_note,
            "run_id": run.id if run else None,
        },
    )
    db.commit()
    return grade_out(db, access.project, assessment)


@router.delete("/grade/{assessment_id}", status_code=204)
def delete_grade(
    assessment_id: int,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    assessment = get_in_project(db, models.GradeAssessment, assessment_id, access.project.id, "GRADE assessment")
    if assessment.status == "signed_off":
        raise HTTPException(
            status_code=409, detail="A signed-off GRADE assessment can't be deleted; edit it to reopen it"
        )
    _audit(db, access, "grade.deleted", "grade_assessment", assessment.id, {"outcome": assessment.outcome})
    db.delete(assessment)
    db.commit()
    return Response(status_code=204)


# --- Summary of findings ---


SOF_COLUMNS = [
    "Outcome",
    "Importance",
    "Studies (participants)",
    "Relative effect (95% CI)",
    "Baseline risk",
    "Risk with comparator",
    "Risk with intervention (95% CI)",
    "Difference (95% CI)",
    "NNT",
    "Certainty",
    "What happens",
    "Explanations",
]


def sof_table_rows(rows: list[dict[str, Any]]) -> list[list[str]]:
    table = []
    for row in rows:
        studies = f"{row['studies']} ({row['participants']})" if row["participants"] else str(row["studies"] or "–")
        absolutes = row["absolute_effects"] or [None]
        for absolute in absolutes:
            table.append(
                [
                    row["outcome"] + (f" ({row['comparison']})" if row["comparison"] else ""),
                    row["importance"].replace("_", " "),
                    studies,
                    row["relative_effect"] or "–",
                    absolute["label"] if absolute else "",
                    f"{absolute['baseline_per_1000']} per 1000" if absolute else "",
                    f"{absolute['intervention_per_1000']} per 1000 ({absolute['intervention_per_1000_ci'][0]} to "
                    f"{absolute['intervention_per_1000_ci'][1]})"
                    if absolute
                    else "",
                    f"{absolute['difference_per_1000']:+d} per 1000 ({absolute['difference_per_1000_ci'][0]:+d} to "
                    f"{absolute['difference_per_1000_ci'][1]:+d})"
                    if absolute
                    else "",
                    f"{absolute['nnt']['type']} {absolute['nnt']['value']} ({absolute['nnt']['interval']})"
                    if absolute and absolute.get("nnt")
                    else "",
                    row["certainty_label"],
                    row["informative_statement"],
                    "; ".join(row["footnotes"]),
                ]
            )
    return table


def summary_of_findings(db: Session, project: models.Project) -> list[dict[str, Any]]:
    order = {"critical": 0, "important": 1, "not_important": 2}
    rows = [sof_row(db, project, a) for a in _assessments(db, project.id)]
    return sorted(rows, key=lambda r: order.get(r["importance"], 3))


@router.get("/summary-of-findings")
def get_summary_of_findings(
    format: Literal["json", "csv", "docx"] = Query("json"),
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    rows = summary_of_findings(db, access.project)
    if format == "json":
        return {
            "intervention": intervention_name(access.project),
            "columns": SOF_COLUMNS,
            "rows": rows,
            "table": sof_table_rows(rows),
        }
    table = sof_table_rows(rows)
    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(SOF_COLUMNS)
        writer.writerows(table)
        return Response(
            buffer.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": 'attachment; filename="summary-of-findings.csv"'},
        )
    document = Document()
    document.add_heading(f"Summary of findings: {access.project.title}", level=1)
    question = access.project.protocol.question if access.project.protocol else ""
    if question:
        document.add_paragraph(question)
    grid = document.add_table(rows=1, cols=len(SOF_COLUMNS))
    grid.style = "Table Grid"
    for cell, label in zip(grid.rows[0].cells, SOF_COLUMNS, strict=True):
        cell.text = label
    for values in table:
        for cell, value in zip(grid.add_row().cells, values, strict=True):
            cell.text = value
    document.add_paragraph(
        "GRADE certainty: High, we are very confident the true effect lies close to the estimate; Moderate, the true "
        "effect is likely close to the estimate but may be substantially different; Low, the true effect may be "
        "substantially different; Very low, the true effect is likely to be substantially different. Risks with the "
        "intervention are based on the assumed comparator risk and the relative effect."
    )
    output = io.BytesIO()
    document.save(output)
    return Response(
        output.getvalue(),
        media_type=DOCX_TYPE,
        headers={"Content-Disposition": 'attachment; filename="summary-of-findings.docx"'},
    )


# --- Evidence to Decision ---


class EtdCriterion(BaseModel):
    judgment: str = Field("", max_length=200)
    research_evidence: str = Field("", max_length=10_000)
    additional_considerations: str = Field("", max_length=10_000)


class EtdIn(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    question: str = Field("", max_length=5000)
    perspective: Literal["clinical_population", "clinical_individual", "health_system"] = "clinical_population"
    criteria: dict[str, EtdCriterion] = Field(default_factory=dict)
    conclusions: dict[str, str] = Field(default_factory=dict)
    grade_assessment_ids: list[int] = Field(default_factory=list, max_length=30)


def etd_out(db: Session, framework: models.EtdFramework) -> dict[str, Any]:
    linked = [db.get(models.GradeAssessment, i) for i in framework.grade_assessment_ids]
    critical = [a for a in linked if a is not None and a.importance == "critical"] or [
        a for a in linked if a is not None
    ]
    lowest = min((grading.LEVELS.index(a.certainty) for a in critical), default=None)
    return {
        "id": framework.id,
        "title": framework.title,
        "question": framework.question,
        "perspective": framework.perspective,
        "criteria": framework.criteria,
        "conclusions": framework.conclusions,
        "grade_assessment_ids": framework.grade_assessment_ids,
        "suggested_certainty": grading.LEVEL_LABELS[grading.LEVELS[lowest]] if lowest is not None else None,
        "status": framework.status,
        "signed_off_at": framework.signed_off_at,
    }


def _apply_etd(db: Session, access: ProjectAccess, framework: models.EtdFramework, body: EtdIn) -> None:
    options = {c["key"]: c["options"] for c in grading.ETD_CRITERIA}
    for key, entry in body.criteria.items():
        if key not in options:
            raise HTTPException(status_code=422, detail=f"Unknown Evidence to Decision criterion: {key}")
        if entry.judgment and entry.judgment not in options[key]:
            raise HTTPException(status_code=422, detail=f"{entry.judgment} isn't a judgment for {key}")
    unknown = set(body.conclusions) - set(grading.CONCLUSION_FIELDS)
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unknown conclusion fields: {', '.join(sorted(unknown))}")
    recommendation_type = body.conclusions.get("recommendation_type", "")
    if recommendation_type and recommendation_type not in grading.RECOMMENDATION_TYPES:
        raise HTTPException(status_code=422, detail="Choose one of the GRADE recommendation types")
    for grade_id in body.grade_assessment_ids:
        get_in_project(db, models.GradeAssessment, grade_id, access.project.id, "GRADE assessment")
    framework.title, framework.question, framework.perspective = body.title.strip(), body.question, body.perspective
    framework.criteria = {k: v.model_dump() for k, v in body.criteria.items()}
    framework.conclusions, framework.grade_assessment_ids = body.conclusions, body.grade_assessment_ids
    framework.status, framework.signed_off_by_id, framework.signed_off_at = "draft", None, None


@router.get("/etd")
def list_etd(access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)):
    frameworks = db.scalars(
        select(models.EtdFramework)
        .where(models.EtdFramework.project_id == access.project.id)
        .order_by(models.EtdFramework.id)
    )
    return [etd_out(db, f) for f in frameworks]


@router.post("/etd", status_code=201)
def create_etd(
    body: EtdIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    framework = models.EtdFramework(project_id=access.project.id, title=body.title, created_by_id=access.user.id)
    _apply_etd(db, access, framework, body)
    db.add(framework)
    db.flush()
    _audit(db, access, "etd.created", "etd_framework", framework.id, {"title": framework.title})
    db.commit()
    return etd_out(db, framework)


@router.put("/etd/{framework_id}")
def update_etd(
    framework_id: int,
    body: EtdIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    framework = get_in_project(
        db, models.EtdFramework, framework_id, access.project.id, "Evidence to Decision framework"
    )
    _apply_etd(db, access, framework, body)
    _audit(
        db,
        access,
        "etd.updated",
        "etd_framework",
        framework.id,
        {"criteria": framework.criteria, "conclusions": framework.conclusions},
    )
    db.commit()
    return etd_out(db, framework)


@router.post("/etd/{framework_id}/sign-off")
def sign_off_etd(
    framework_id: int,
    body: SignOffIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    framework = get_in_project(
        db, models.EtdFramework, framework_id, access.project.id, "Evidence to Decision framework"
    )
    missing = [c["label"] for c in grading.ETD_CRITERIA if not (framework.criteria.get(c["key"]) or {}).get("judgment")]
    if missing:
        raise HTTPException(status_code=409, detail=f"Make a judgment on every criterion first: {'; '.join(missing)}")
    if (
        not framework.conclusions.get("recommendation_type")
        or not framework.conclusions.get("justification", "").strip()
    ):
        raise HTTPException(status_code=409, detail="Choose the type of recommendation and justify it")
    unsigned = [
        i
        for i in framework.grade_assessment_ids
        if (a := db.get(models.GradeAssessment, i)) is None or a.status != "signed_off"
    ]
    if unsigned:
        raise HTTPException(status_code=409, detail="Sign off the linked GRADE assessments first")
    framework.status, framework.signed_off_by_id, framework.signed_off_at = (
        "signed_off",
        access.user.id,
        models.utcnow(),
    )
    _audit(
        db,
        access,
        "etd.signed_off",
        "etd_framework",
        framework.id,
        {"note": body.note, "conclusions": framework.conclusions},
    )
    db.commit()
    return etd_out(db, framework)


# --- Prior reviews ---


class PriorReviewIn(BaseModel):
    title: str = Field(min_length=1, max_length=2000)
    doi: str = Field("", max_length=255)
    year: str = Field("", max_length=20)
    openalex_id: str = Field("", max_length=40)
    outcome: str = Field("", max_length=300)
    conclusion: str = Field("", max_length=10_000)
    conclusion_direction: Literal["", "favours_intervention", "favours_comparator", "no_difference", "uncertain"] = ""
    fetch_references: bool = True


def prior_out(review: models.PriorReview) -> dict[str, Any]:
    return {
        "id": review.id,
        "title": review.title,
        "doi": review.doi,
        "year": review.year,
        "openalex_id": review.openalex_id,
        "references": len(review.referenced_works),
        "outcome": review.outcome,
        "conclusion": review.conclusion,
        "conclusion_direction": review.conclusion_direction,
    }


def fetch_review_work(doi: str, openalex_id: str) -> dict[str, Any] | None:
    select_fields = "id,doi,title,publication_year,referenced_works"
    if openalex_id:
        works = openalex_raw_works("openalex", [openalex_id], select_fields)
    elif doi:
        works = openalex_raw_works("doi", [doi.lower()], select_fields)
    else:
        return None
    return works[0] if works else None


def _tail(url: str) -> str:
    return url.rstrip("/").split("/")[-1]


@router.get("/prior-reviews")
def list_prior_reviews(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    reviews = db.scalars(
        select(models.PriorReview)
        .where(models.PriorReview.project_id == access.project.id)
        .order_by(models.PriorReview.id)
    )
    latest = db.scalar(
        select(models.TopicExploration)
        .where(models.TopicExploration.project_id == access.project.id)
        .order_by(models.TopicExploration.id.desc())
        .limit(1)
    )
    return {
        "reviews": [prior_out(r) for r in reviews],
        "from_topic_exploration": (latest.results.get("existing_reviews", []) if latest else []),
    }


@router.post("/prior-reviews", status_code=201)
def add_prior_review(
    body: PriorReviewIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    review = models.PriorReview(
        project_id=access.project.id,
        title=body.title.strip(),
        doi=body.doi.strip().lower().removeprefix("https://doi.org/"),
        year=body.year,
        openalex_id=body.openalex_id.strip(),
        outcome=body.outcome.strip(),
        conclusion=body.conclusion.strip(),
        conclusion_direction=body.conclusion_direction,
        created_by_id=access.user.id,
    )
    warning = None
    if body.fetch_references and (review.doi or review.openalex_id):
        try:
            work = fetch_review_work(review.doi, review.openalex_id)
        except SearchError:
            work, warning = None, "OpenAlex couldn't be reached, so the review's references weren't loaded."
        if work is not None:
            review.openalex_id = _tail(work.get("id") or "")
            review.referenced_works = [_tail(w) for w in work.get("referenced_works") or []]
            review.year = review.year or str(work.get("publication_year") or "")
        elif warning is None:
            warning = "OpenAlex has no record of this review, so its references weren't loaded."
    db.add(review)
    db.flush()
    _audit(
        db,
        access,
        "prior_review.added",
        "prior_review",
        review.id,
        {"title": review.title, "doi": review.doi, "references": len(review.referenced_works)},
    )
    db.commit()
    return {**prior_out(review), "warning": warning}


@router.put("/prior-reviews/{review_id}")
def update_prior_review(
    review_id: int,
    body: PriorReviewIn,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    review = get_in_project(db, models.PriorReview, review_id, access.project.id, "Prior review")
    review.title, review.year, review.outcome = body.title.strip(), body.year, body.outcome.strip()
    review.conclusion, review.conclusion_direction = body.conclusion.strip(), body.conclusion_direction
    _audit(
        db,
        access,
        "prior_review.updated",
        "prior_review",
        review.id,
        {"conclusion_direction": review.conclusion_direction},
    )
    db.commit()
    return prior_out(review)


@router.delete("/prior-reviews/{review_id}", status_code=204)
def delete_prior_review(
    review_id: int,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    review = get_in_project(db, models.PriorReview, review_id, access.project.id, "Prior review")
    _audit(db, access, "prior_review.deleted", "prior_review", review.id, {"title": review.title})
    db.delete(review)
    db.commit()
    return Response(status_code=204)


def our_direction(row: dict[str, Any], outcome_direction: str) -> str:
    if row["certainty"] == "very_low" or row["direction"] == "unknown":
        return "uncertain"
    if row["direction"] == "little_or_no_difference":
        return "no_difference"
    lower_is_better = outcome_direction == "lower_is_better"
    return "favours_intervention" if (row["direction"] == "reduction") == lower_is_better else "favours_comparator"


@router.get("/prior-reviews/comparison")
def prior_review_comparison(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    """Which included studies each prior review cites (overlap matrix and corrected covered area), and conclusions that
    differ from this review's for the same outcome."""
    reviews = list(
        db.scalars(
            select(models.PriorReview)
            .where(models.PriorReview.project_id == access.project.id)
            .order_by(models.PriorReview.id)
        )
    )
    studies = included_studies(db, access.project.id)
    records = [report.record for study in studies for report in study.reports]
    try:
        openalex_ids = resolve_openalex_ids(records) if records and reviews else {}
        lookup_error = None
    except SearchError:
        openalex_ids, lookup_error = {}, "OpenAlex couldn't be reached to identify the included studies."
    matrix: list[dict[str, Any]] = []
    for study in studies:
        ids = {openalex_ids[r.record_id] for r in study.reports if r.record_id in openalex_ids}
        matrix.append(
            {
                "study_id": study.id,
                "label": study.label,
                "identified": bool(ids),
                "cited_by": [bool(ids & set(review.referenced_works)) for review in reviews],
            }
        )
    boolean = [[True, *row["cited_by"]] for row in matrix if row["identified"]]
    cca = grading.corrected_covered_area(boolean) if reviews else None
    grades = {a.outcome.casefold(): a for a in _assessments(db, access.project.id)}
    conclusions = []
    for review in reviews:
        assessment = grades.get(review.outcome.casefold()) if review.outcome else None
        if assessment is None or not review.conclusion_direction:
            continue
        ours = our_direction(sof_row(db, access.project, assessment), assessment.outcome_direction)
        conclusions.append(
            {
                "prior_review_id": review.id,
                "outcome": review.outcome,
                "prior": review.conclusion_direction,
                "this_review": ours,
                "changed": ours != review.conclusion_direction,
            }
        )
    return {
        "reviews": [prior_out(r) for r in reviews],
        "matrix": matrix,
        "corrected_covered_area": cca,
        "overlap": grading.overlap_label(cca),
        "note": (
            "The matrix covers this review's included studies found in OpenAlex; the corrected covered area counts "
            "this review as one column."
        ),
        "lookup_error": lookup_error,
        "conclusions": conclusions,
    }


# --- Interpretation ---


NUMBER = re.compile(r"(?<![\w.])-?\d{1,3}(?:,\d{3})+(?:\.\d+)?|(?<![\w.])-?\d+(?:\.\d+)?")


def _numbers(text: str) -> list[str]:
    return NUMBER.findall(text.replace("−", "-"))


def unverified_numbers(text: str, allowed_texts: list[str]) -> list[str]:
    """Numbers in the text that don't appear in the allowed texts (the summary of findings as displayed), allowing for
    rounding and for per 1000 figures written as percentages."""
    allowed = {95.0, 100.0, 1000.0}
    for allowed_text in allowed_texts:
        allowed |= {float(token.replace(",", "")) for token in _numbers(allowed_text)}
    candidates = allowed | {value / 10 for value in allowed}
    problems = []
    for token in _numbers(text):
        value = float(token.replace(",", ""))
        decimals = len(token.split(".")[1]) if "." in token else 0
        if not any(round(abs(candidate), decimals) == abs(value) for candidate in candidates):
            problems.append(token)
    return sorted(set(problems))


def checked_numbers(db: Session, project: models.Project, text: models.InterpretationText) -> list[str]:
    """Only AI-written plain-language summaries are checked; rules-based text is built from the results themselves."""
    if text.kind != "plain_language_summary":
        return []
    return unverified_numbers(
        text.content, [" ".join(values) for values in sof_table_rows(summary_of_findings(db, project))]
    )


def interpretation_out(text: models.InterpretationText) -> dict[str, Any]:
    return {
        "id": text.id,
        "kind": text.kind,
        "outcome": text.outcome,
        "content": text.content,
        "generated_by": text.generated_by,
        "ai_run_id": text.ai_run_id,
        "unverified_numbers": text.unverified_numbers,
        "status": text.status,
        "approved_at": text.approved_at,
        "updated_at": text.updated_at,
    }


def _texts(db: Session, project_id: int) -> list[models.InterpretationText]:
    return list(
        db.scalars(
            select(models.InterpretationText)
            .where(models.InterpretationText.project_id == project_id)
            .order_by(models.InterpretationText.id)
        )
    )


@router.get("/interpretation")
def list_interpretation(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    return [interpretation_out(t) for t in _texts(db, access.project.id)]


class RulesRequest(BaseModel):
    kind: Literal["informative_statement", "limitations"]


@router.post("/interpretation/rules", status_code=201)
def generate_rules_text(
    body: RulesRequest,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    """Informative statements (GRADE wording) or limitations from the GRADE ratings and analysis notes. Replaces
    earlier drafts of the same kind; approved texts are kept."""
    require_stage_open(db, access.project.id, "certainty")
    assessments = _assessments(db, access.project.id)
    if not assessments:
        raise HTTPException(status_code=409, detail="Create GRADE assessments first")
    for text in _texts(db, access.project.id):
        if text.kind == body.kind and text.status == "draft" and text.generated_by == "rules":
            db.delete(text)
    created = []
    if body.kind == "informative_statement":
        for assessment in assessments:
            row = sof_row(db, access.project, assessment)
            created.append(
                models.InterpretationText(
                    project_id=access.project.id,
                    kind=body.kind,
                    outcome=assessment.outcome,
                    content=row["informative_statement"],
                    generated_by="rules",
                    status="draft",
                    created_by_id=access.user.id,
                )
            )
    else:
        notes = []
        for assessment in assessments:
            analysis, run = analysis_for(db, assessment)
            if run is not None and analysis is not None:
                notes += [f"{analysis.title}: {note}" for note in (run.results or {}).get("notes", [])]
                excluded = run.dataset.get("excluded", [])
                if excluded:
                    notes.append(
                        f"{analysis.title}: {len(excluded)} studies couldn't contribute data "
                        f"({'; '.join(sorted({e['reason'] for e in excluded}))})."
                    )
        statements = grading.limitations([{"outcome": a.outcome, "domains": a.domains} for a in assessments], notes)
        content = (
            "\n".join(f"- {s}" for s in statements) or "- No limitations were flagged by the GRADE ratings or analyses."
        )
        created.append(
            models.InterpretationText(
                project_id=access.project.id,
                kind=body.kind,
                content=content,
                generated_by="rules",
                status="draft",
                created_by_id=access.user.id,
            )
        )
    db.add_all(created)
    db.flush()
    _audit(
        db,
        access,
        "interpretation.generated",
        "project",
        access.project.id,
        {"kind": body.kind, "texts": [t.id for t in created]},
    )
    db.commit()
    return [interpretation_out(t) for t in created]


class PlainLanguageOutput(BaseModel):
    summary: str


@router.post("/interpretation/plain-language", status_code=201, dependencies=[Depends(ai_rate_limit)])
async def generate_plain_language(
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)), db: Session = Depends(get_db)
):
    require_stage_open(db, access.project.id, "certainty")
    rows = summary_of_findings(db, access.project)
    if not rows:
        raise HTTPException(status_code=409, detail="Create GRADE assessments first")
    findings = "\n".join(f"- {' | '.join(values)}" for values in sof_table_rows(rows))
    question = (access.project.protocol.question if access.project.protocol else "") or access.project.title
    ai = project_ai(db, access)
    run = new_ai_run(access, "plain_language", PLAIN_LANGUAGE_PROMPT, ai)
    prompt = PLAIN_LANGUAGE_PROMPT.render(question=question, findings=" | ".join(SOF_COLUMNS) + "\n" + findings)
    try:
        result = await complete_structured(ai, prompt, PlainLanguageOutput, max_tokens=2000)
    except LLMError as exc:
        run.status, run.error = "failed", str(exc)
        record_usage(run, ai, exc.usage)
        db.add(run)
        db.commit()
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    run.status = "succeeded"
    record_usage(run, ai, result.usage)
    db.add(run)
    db.flush()
    content = result.value.summary.strip()
    text = models.InterpretationText(
        project_id=access.project.id,
        kind="plain_language_summary",
        content=content,
        generated_by="ai",
        ai_run_id=run.id,
        unverified_numbers=unverified_numbers(content, [" ".join(values) for values in sof_table_rows(rows)]),
        status="draft",
        created_by_id=access.user.id,
    )
    db.add(text)
    db.flush()
    _audit(
        db,
        access,
        "ai.plain_language_summary",
        "interpretation_text",
        text.id,
        {
            "provider": run.provider,
            "model": run.model,
            "prompt_version": run.prompt_version,
            "unverified_numbers": text.unverified_numbers,
        },
    )
    db.commit()
    return interpretation_out(text)


class InterpretationUpdate(BaseModel):
    content: str = Field(min_length=1, max_length=20_000)


@router.put("/interpretation/{text_id}")
def update_interpretation(
    text_id: int,
    body: InterpretationUpdate,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    text = get_in_project(db, models.InterpretationText, text_id, access.project.id, "Interpretation text")
    before = text.content
    text.content = body.content.strip()
    text.unverified_numbers = checked_numbers(db, access.project, text)
    text.status, text.approved_by_id, text.approved_at = "draft", None, None
    _audit(
        db, access, "interpretation.edited", "interpretation_text", text.id, {"before": before, "after": text.content}
    )
    db.commit()
    return interpretation_out(text)


@router.post("/interpretation/{text_id}/approve")
def approve_interpretation(
    text_id: int,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_INTERPRETATION)),
    db: Session = Depends(get_db),
):
    """A clinical expert (or methodologist) approves interpretive text. Numbers must match the summary of findings."""
    require_stage_open(db, access.project.id, "certainty")
    text = get_in_project(db, models.InterpretationText, text_id, access.project.id, "Interpretation text")
    text.unverified_numbers = checked_numbers(db, access.project, text)
    if text.unverified_numbers:
        db.commit()
        raise HTTPException(
            status_code=409,
            detail=f"These numbers aren't in the summary of findings: {', '.join(text.unverified_numbers)}. Correct "
            "them before approving.",
        )
    text.status, text.approved_by_id, text.approved_at = "approved", access.user.id, models.utcnow()
    _audit(
        db,
        access,
        "interpretation.approved",
        "interpretation_text",
        text.id,
        {"kind": text.kind, "content": text.content},
    )
    db.commit()
    return interpretation_out(text)


@router.delete("/interpretation/{text_id}", status_code=204)
def delete_interpretation(
    text_id: int,
    access: ProjectAccess = Depends(project_access(Permission.APPROVE_CERTAINTY)),
    db: Session = Depends(get_db),
):
    require_stage_open(db, access.project.id, "certainty")
    text = get_in_project(db, models.InterpretationText, text_id, access.project.id, "Interpretation text")
    _audit(db, access, "interpretation.deleted", "interpretation_text", text.id, {"kind": text.kind})
    db.delete(text)
    db.commit()
    return Response(status_code=204)
