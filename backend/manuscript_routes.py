"""The manuscript: sections drafted from the locked evidence base, claim verification, references and citation styles,
reporting checklists, tables and figures, AI drafts and language edits (as suggestions), versions approved by every
author, and export."""

import asyncio
import csv
import difflib
import io
import json
import re
from datetime import timedelta
from typing import Literal

from docx import Document
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

import citations
import models
from ai_access import new_ai_run, project_ai, record_usage
from audit import record_event
from claims import ACKNOWLEDGEABLE, check_content, sentence_hash, split_sentences
from database import get_db
from evidence_catalog import build_catalog, latest_snapshot
from extraction_data import included_studies
from llm.prompts import LANGUAGE_EDIT_PROMPT, MANUSCRIPT_DRAFT_PROMPT, MANUSCRIPT_EXTRAS_PROMPT
from manuscript_assets import available_assets, figure, table
from manuscript_checklists import CHECKLISTS, STATUSES, applicable, checklist_context, complete, evaluate
from manuscript_content import SECTION_SPECS, SECTIONS, generate
from manuscript_export import ExportUnavailable, export_manuscript
from manuscript_state import (
    CREDIT_ROLES,
    STATEMENT_KEYS,
    approval_state,
    cited_ids,
    content_sha,
    get_manuscript,
    latest_version,
    reference_states,
    references,
    verify,
    version_content,
)
from permissions import Permission
from projects_routes import ProjectAccess, get_in_project, project_access
from protocol_frameworks import PROTOCOL_SECTIONS
from publishing_tools import tools_status
from rate_limiting import ai_rate_limit
from services.ai_manuscript import (
    LANGUAGE_TASKS,
    assemble_draft,
    draft_section,
    edit_language,
    preservation_problems,
    unverified_numbers,
    write_extra,
)
from services.errors import LLMError
from services.reference_checks import (
    ReferenceCheckError,
    check_reference,
    crossref_work,
    pubmed_summary,
    zotero_create,
    zotero_items,
)
from workflow import require_stage_open

router = APIRouter(prefix="/api/projects/{project_id}", tags=["manuscript"])

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
RECHECK_AFTER = timedelta(days=7)


def _audit(db: Session, access: ProjectAccess, action: str, entity_type: str, entity_id: int, details: dict) -> None:
    record_event(
        db,
        project_id=access.project.id,
        actor_id=access.user.id,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        details=details,
    )


def _manuscript(db: Session, access: ProjectAccess) -> models.Manuscript:
    manuscript = get_manuscript(db, access.project.id)
    if manuscript is None:
        raise HTTPException(status_code=404, detail="Start the manuscript first")
    return manuscript


def _section(manuscript: models.Manuscript, key: str) -> models.ManuscriptSection:
    section = next((s for s in manuscript.sections if s.key == key), None)
    if section is None:
        raise HTTPException(status_code=404, detail=f"The manuscript has no section {key}")
    return section


def _open(db: Session, access: ProjectAccess) -> None:
    require_stage_open(db, access.project.id, "manuscript")


def _save_section(
    db: Session,
    access: ProjectAccess,
    section: models.ManuscriptSection,
    content: str,
    source: str,
    note: str,
    suggestion_id: int | None = None,
) -> models.ManuscriptRevision:
    section.content = content
    section.updated_by_id = access.user.id
    revision = models.ManuscriptRevision(
        section_id=section.id,
        content=content,
        source=source,
        note=note,
        suggestion_id=suggestion_id,
        created_by_id=access.user.id,
    )
    db.add(revision)
    db.flush()
    return revision


def reference_out(ref: models.ManuscriptReference, style: str, cited: set[int]) -> dict:
    return {
        "id": ref.id,
        "record_id": ref.record_id,
        "csl": ref.csl,
        "doi": ref.doi,
        "pmid": ref.pmid,
        "formatted": citations.format_reference(ref.csl, style if style in citations.BUILT_IN_STYLES else "vancouver"),
        "cited": ref.id in cited,
        "verification_status": ref.verification_status,
        "verification": ref.verification,
        "retracted": ref.retracted,
        "retraction": ref.retraction,
        "checked_at": ref.checked_at,
    }


def manuscript_out(db: Session, project: models.Project, manuscript: models.Manuscript) -> dict:
    return {
        "id": manuscript.id,
        "title": manuscript.title,
        "citation_style": manuscript.citation_style,
        "built_in_styles": citations.BUILT_IN_STYLES,
        "keywords": manuscript.keywords,
        "statements": manuscript.statements,
        "extras": manuscript.extras,
        "evidence_snapshot_id": manuscript.evidence_snapshot_id,
        "sections": [
            {
                "key": s.key,
                "title": s.title,
                "content": s.content,
                "updated_at": s.updated_at,
                "guidance": SECTION_SPECS[s.key].guidance if s.key in SECTION_SPECS else "",
            }
            for s in sorted(manuscript.sections, key=lambda s: s.position)
        ],
        "authors": [
            {
                "id": a.id,
                "user_id": a.user_id,
                "name": a.name,
                "email": a.email,
                "affiliation": a.affiliation,
                "orcid": a.orcid,
                "corresponding": a.corresponding,
                "credit_roles": a.credit_roles,
                "competing_interests": a.competing_interests,
            }
            for a in manuscript.authors
        ],
        "credit_roles": CREDIT_ROLES,
        "statement_keys": STATEMENT_KEYS,
        "verification": verify(db, project, manuscript),
        "approvals": approval_state(db, manuscript),
        "versions": [
            {"id": v.id, "number": v.number, "note": v.note, "created_at": v.created_at, "sha256": v.sha256}
            for v in manuscript.versions
        ],
        "tools": tools_status(),
        "updated_at": manuscript.updated_at,
    }


@router.get("/manuscript")
def get_manuscript_route(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    manuscript = get_manuscript(db, access.project.id)
    return manuscript_out(db, access.project, manuscript) if manuscript else None


class ManuscriptCreate(BaseModel):
    title: str = Field("", max_length=500)
    citation_style: str = Field("vancouver", max_length=100)


def _check_style(style: str) -> None:
    if style not in citations.BUILT_IN_STYLES and not citations.CSL_STYLE_ID.match(style):
        raise HTTPException(status_code=422, detail="Choose a built-in citation style or a CSL style id such as nature")


@router.post("/manuscript", status_code=201)
def create_manuscript(
    body: ManuscriptCreate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Start the manuscript from the locked evidence base: PRISMA sections, text generated from the review's records,
    and references for the included studies."""
    _open(db, access)
    if get_manuscript(db, access.project.id) is not None:
        raise HTTPException(status_code=409, detail="This project already has a manuscript")
    _check_style(body.citation_style)
    project = access.project
    question = project.protocol.question if project.protocol else ""
    snapshot = latest_snapshot(db, project.id, "certainty")
    manuscript = models.Manuscript(
        project_id=project.id,
        title=body.title.strip() or f"{project.title}: a systematic review",
        citation_style=body.citation_style,
        keywords=[],
        statements={},
        extras={},
        evidence_snapshot_id=snapshot.id if snapshot else None,
        created_by_id=access.user.id,
    )
    db.add(manuscript)
    db.flush()
    study_references: dict[int, list[int]] = {}
    for study in included_studies(db, project.id):
        for report in study.reports:
            ref = models.ManuscriptReference(
                project_id=project.id,
                record_id=report.record_id,
                csl={},
                doi=report.record.doi,
                pmid=(report.record.identifiers or {}).get("pmid", ""),
                created_by_id=access.user.id,
            )
            db.add(ref)
            db.flush()
            ref.csl = citations.csl_from_record(ref.id, report.record)
            study_references.setdefault(study.id, []).append(ref.id)
    for position, spec in enumerate(SECTIONS):
        section = models.ManuscriptSection(
            manuscript_id=manuscript.id,
            key=spec.key,
            title=spec.title,
            position=position,
            content="",
            updated_by_id=access.user.id,
        )
        db.add(section)
        db.flush()
        manuscript.sections.append(section)
        content = generate(db, project, manuscript, spec.key, study_references) if spec.generated else spec.skeleton
        _save_section(db, access, section, content, "generated", "Started from the locked evidence base")
    _audit(
        db,
        access,
        "manuscript.created",
        "manuscript",
        manuscript.id,
        {"title": manuscript.title, "evidence_snapshot_id": manuscript.evidence_snapshot_id, "question": question},
    )
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, project, manuscript)


class ManuscriptUpdate(BaseModel):
    title: str = Field(min_length=1, max_length=500)
    citation_style: str = Field(max_length=100)
    keywords: list[str] = Field(default_factory=list, max_length=20)
    statements: dict[str, str] = Field(default_factory=dict)
    extras: dict[str, list[str]] = Field(default_factory=dict)


@router.patch("/manuscript")
def update_manuscript(
    body: ManuscriptUpdate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    _check_style(body.citation_style)
    unknown = set(body.statements) - set(STATEMENT_KEYS)
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unknown statements: {', '.join(sorted(unknown))}")
    before = {
        "title": manuscript.title,
        "citation_style": manuscript.citation_style,
        "statements": manuscript.statements,
    }
    manuscript.title, manuscript.citation_style = body.title.strip(), body.citation_style
    manuscript.keywords = [k.strip() for k in body.keywords if k.strip()]
    manuscript.statements = {k: v.strip() for k, v in body.statements.items()}
    extras = dict(manuscript.extras)
    for kind, items in body.extras.items():
        if kind in ("graphical_abstract", "highlights"):
            extras[kind] = {
                **(extras.get(kind) or {}),
                "items": [i.strip() for i in items if i.strip()],
                "edited": True,
            }
    manuscript.extras = extras
    _audit(
        db,
        access,
        "manuscript.updated",
        "manuscript",
        manuscript.id,
        {
            "before": before,
            "after": {
                "title": manuscript.title,
                "citation_style": manuscript.citation_style,
                "statements": manuscript.statements,
            },
        },
    )
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


@router.get("/manuscript/evidence")
def manuscript_evidence(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    manuscript = get_manuscript(db, access.project.id)
    catalog = build_catalog(db, access.project, manuscript)
    return [
        {"key": item.key, "marker": f"[#{item.key}]", "label": item.label, "group": item.group, "facts": item.facts}
        for item in catalog.values()
    ]


class SectionUpdate(BaseModel):
    content: str = Field(max_length=200_000)
    note: str = Field("", max_length=2000)


@router.put("/manuscript/sections/{key}")
def save_section(
    key: str,
    body: SectionUpdate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, key)
    if section.content != body.content:
        revision = _save_section(db, access, section, body.content, "manual", body.note)
        _audit(
            db,
            access,
            "manuscript.section_saved",
            "manuscript_section",
            section.id,
            {"key": key, "revision_id": revision.id, "characters": len(body.content)},
        )
        db.commit()
        db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


def _word_diff(before: str, after: str) -> list[dict]:
    a, b = re.findall(r"\S+|\s+", before), re.findall(r"\S+|\s+", after)
    ops = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
        if tag == "equal":
            ops.append({"op": "equal", "text": "".join(a[i1:i2])})
            continue
        if i2 > i1:
            ops.append({"op": "delete", "text": "".join(a[i1:i2])})
        if j2 > j1:
            ops.append({"op": "insert", "text": "".join(b[j1:j2])})
    return ops


@router.get("/manuscript/sections/{key}/revisions")
def section_revisions(
    key: str, access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    section = _section(_manuscript(db, access), key)
    revisions = db.scalars(
        select(models.ManuscriptRevision)
        .where(models.ManuscriptRevision.section_id == section.id)
        .order_by(models.ManuscriptRevision.id.desc())
    )
    return [
        {
            "id": r.id,
            "source": r.source,
            "note": r.note,
            "created_by": r.created_by.full_name if r.created_by else None,
            "created_at": r.created_at,
            "characters": len(r.content),
        }
        for r in revisions
    ]


@router.get("/manuscript/revisions/{revision_id}/diff")
def revision_diff(
    revision_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """Tracked changes: a revision compared with the revision before it."""
    manuscript = _manuscript(db, access)
    revision = db.get(models.ManuscriptRevision, revision_id)
    if revision is None or revision.section_id not in {s.id for s in manuscript.sections}:
        raise HTTPException(status_code=404, detail="Revision not found")
    previous = db.scalar(
        select(models.ManuscriptRevision)
        .where(models.ManuscriptRevision.section_id == revision.section_id, models.ManuscriptRevision.id < revision.id)
        .order_by(models.ManuscriptRevision.id.desc())
        .limit(1)
    )
    return {
        "revision_id": revision.id,
        "previous_id": previous.id if previous else None,
        "diff": _word_diff(previous.content if previous else "", revision.content),
    }


@router.post("/manuscript/revisions/{revision_id}/restore")
def restore_revision(
    revision_id: int,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    revision = db.get(models.ManuscriptRevision, revision_id)
    section = next((s for s in manuscript.sections if revision and s.id == revision.section_id), None)
    if revision is None or section is None:
        raise HTTPException(status_code=404, detail="Revision not found")
    _save_section(db, access, section, revision.content, "restored", f"Restored revision {revision.id}")
    _audit(db, access, "manuscript.revision_restored", "manuscript_section", section.id, {"revision_id": revision.id})
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


def suggestion_out(
    db: Session,
    project: models.Project,
    manuscript: models.Manuscript,
    suggestion: models.ManuscriptSuggestion,
    catalog=None,
) -> dict:
    section = next(s for s in manuscript.sections if s.id == suggestion.section_id)
    catalog = catalog if catalog is not None else build_catalog(db, project, manuscript)
    checks = check_content(suggestion.content, catalog, reference_states(db, project.id), set())
    return {
        "id": suggestion.id,
        "section_key": section.key,
        "kind": suggestion.kind,
        "instruction": suggestion.instruction,
        "content": suggestion.content,
        "problems": suggestion.problems,
        "status": suggestion.status,
        "ai_run_id": suggestion.ai_run_id,
        "created_at": suggestion.created_at,
        "diff": _word_diff(section.content, suggestion.content),
        "verification": {
            "sentences": [c.out() for c in checks],
            "unverified": sum(1 for c in checks if c.status != "verified"),
        },
    }


def _new_suggestion(
    db: Session,
    access: ProjectAccess,
    section: models.ManuscriptSection,
    kind: str,
    content: str,
    problems: list[str],
    run: models.AIRun | None = None,
    instruction: str = "",
) -> models.ManuscriptSuggestion:
    suggestion = models.ManuscriptSuggestion(
        section_id=section.id,
        kind=kind,
        instruction=instruction,
        original=section.content,
        content=content,
        problems=problems,
        ai_run_id=run.id if run else None,
        status="pending",
        created_by_id=access.user.id,
    )
    db.add(suggestion)
    db.flush()
    return suggestion


@router.post("/manuscript/sections/{key}/regenerate", status_code=201)
def regenerate_section(
    key: str, access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)), db: Session = Depends(get_db)
):
    """Rebuild a generated section from the current records, as a suggestion to review against the edited text."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, key)
    if not SECTION_SPECS[key].generated:
        raise HTTPException(
            status_code=422, detail="Only methods, results, discussion, other information, and AI use are generated"
        )
    study_references: dict[int, list[int]] = {}
    for ref in references(db, access.project.id):
        if ref.record_id is None:
            continue
        report = db.scalar(select(models.StudyReport).where(models.StudyReport.record_id == ref.record_id))
        if report is not None:
            study_references.setdefault(report.study_id, []).append(ref.id)
    suggestion = _new_suggestion(
        db, access, section, "generated", generate(db, access.project, manuscript, key, study_references), []
    )
    _audit(
        db,
        access,
        "manuscript.section_regenerated",
        "manuscript_section",
        section.id,
        {"key": key, "suggestion_id": suggestion.id},
    )
    db.commit()
    return suggestion_out(db, access.project, manuscript, suggestion)


def _review_summary(project: models.Project) -> str:
    protocol = project.protocol
    return f"Title: {project.title}\nQuestion: {protocol.question if protocol else ''}"


@router.post("/manuscript/sections/{key}/draft", status_code=201, dependencies=[Depends(ai_rate_limit)])
async def draft_with_ai(
    key: str, access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)), db: Session = Depends(get_db)
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, key)
    catalog = build_catalog(db, access.project, manuscript)
    refs = {r.id: r.csl for r in references(db, access.project.id) if not r.retracted}
    ai = project_ai(db, access)
    run = new_ai_run(access, "manuscript", MANUSCRIPT_DRAFT_PROMPT, ai)
    try:
        result = await draft_section(
            ai,
            section.title,
            SECTION_SPECS[key].guidance,
            _review_summary(access.project),
            catalog,
            refs,
            section.content,
        )
    except LLMError as exc:
        run.status, run.error = "failed", str(exc)
        record_usage(run, ai, exc.usage)
        db.add(run)
        db.commit()
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    run.status = "succeeded"
    record_usage(run, ai, result.usage)
    db.add(run)
    db.flush()
    content, problems = assemble_draft(result.value, catalog, refs)
    suggestion = _new_suggestion(db, access, section, "draft", content, problems, run)
    _audit(
        db,
        access,
        "ai.manuscript_draft",
        "manuscript_section",
        section.id,
        {
            "key": key,
            "suggestion_id": suggestion.id,
            "provider": run.provider,
            "model": run.model,
            "prompt_version": run.prompt_version,
        },
    )
    db.commit()
    return suggestion_out(db, access.project, manuscript, suggestion, catalog)


class LanguageRequest(BaseModel):
    kind: Literal["academic_tone", "grammar", "journal_style", "plain_language", "concise"]
    journal_style: str = Field("", max_length=5000)


@router.post("/manuscript/sections/{key}/language", status_code=201, dependencies=[Depends(ai_rate_limit)])
async def language_edit(
    key: str,
    body: LanguageRequest,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, key)
    if not section.content.strip():
        raise HTTPException(status_code=422, detail="Write the section before editing its language")
    ai = project_ai(db, access)
    run = new_ai_run(access, "manuscript", LANGUAGE_EDIT_PROMPT, ai)
    try:
        result = await edit_language(ai, body.kind, section.content, body.journal_style)
    except LLMError as exc:
        run.status, run.error = "failed", str(exc)
        record_usage(run, ai, exc.usage)
        db.add(run)
        db.commit()
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    run.status = "succeeded"
    record_usage(run, ai, result.usage)
    db.add(run)
    db.flush()
    edited = result.value.strip() + "\n"
    suggestion = _new_suggestion(
        db,
        access,
        section,
        body.kind,
        edited,
        preservation_problems(section.content, edited),
        run,
        LANGUAGE_TASKS[body.kind],
    )
    _audit(
        db,
        access,
        "ai.manuscript_language",
        "manuscript_section",
        section.id,
        {"key": key, "kind": body.kind, "suggestion_id": suggestion.id, "problems": suggestion.problems},
    )
    db.commit()
    return suggestion_out(db, access.project, manuscript, suggestion)


@router.get("/manuscript/suggestions")
def list_suggestions(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    manuscript = _manuscript(db, access)
    rows = db.scalars(
        select(models.ManuscriptSuggestion)
        .where(
            models.ManuscriptSuggestion.section_id.in_([s.id for s in manuscript.sections]),
            models.ManuscriptSuggestion.status == "pending",
        )
        .order_by(models.ManuscriptSuggestion.id.desc())
    ).all()
    catalog = build_catalog(db, access.project, manuscript) if rows else {}
    return [suggestion_out(db, access.project, manuscript, s, catalog) for s in rows]


class SuggestionDecision(BaseModel):
    acknowledge_problems: bool = False


def _suggestion(
    db: Session, manuscript: models.Manuscript, suggestion_id: int
) -> tuple[models.ManuscriptSuggestion, models.ManuscriptSection]:
    suggestion = db.get(models.ManuscriptSuggestion, suggestion_id)
    section = next((s for s in manuscript.sections if suggestion and s.id == suggestion.section_id), None)
    if suggestion is None or section is None:
        raise HTTPException(status_code=404, detail="Suggestion not found")
    if suggestion.status != "pending":
        raise HTTPException(status_code=409, detail="This suggestion was already decided")
    return suggestion, section


@router.post("/manuscript/suggestions/{suggestion_id}/accept")
def accept_suggestion(
    suggestion_id: int,
    body: SuggestionDecision,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    suggestion, section = _suggestion(db, manuscript, suggestion_id)
    if suggestion.problems and not body.acknowledge_problems:
        raise HTTPException(
            status_code=409, detail=f"Review the problems before accepting: {'; '.join(suggestion.problems)}"
        )
    if section.content != suggestion.original:
        raise HTTPException(
            status_code=409, detail="The section changed after this suggestion was made; ask for a new one"
        )
    source = "generated" if suggestion.kind == "generated" else "ai_accepted"
    _save_section(
        db,
        access,
        section,
        suggestion.content,
        source,
        f"Accepted {suggestion.kind.replace('_', ' ')} suggestion",
        suggestion.id,
    )
    suggestion.status, suggestion.decided_by_id, suggestion.decided_at = "accepted", access.user.id, models.utcnow()
    _audit(
        db,
        access,
        "manuscript.suggestion_accepted",
        "manuscript_suggestion",
        suggestion.id,
        {"key": section.key, "kind": suggestion.kind, "problems": suggestion.problems},
    )
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


@router.post("/manuscript/suggestions/{suggestion_id}/reject")
def reject_suggestion(
    suggestion_id: int,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    manuscript = _manuscript(db, access)
    suggestion, section = _suggestion(db, manuscript, suggestion_id)
    suggestion.status, suggestion.decided_by_id, suggestion.decided_at = "rejected", access.user.id, models.utcnow()
    _audit(
        db,
        access,
        "manuscript.suggestion_rejected",
        "manuscript_suggestion",
        suggestion.id,
        {"key": section.key, "kind": suggestion.kind},
    )
    db.commit()
    return {"id": suggestion.id, "status": suggestion.status}


@router.post("/manuscript/extras/{kind}", dependencies=[Depends(ai_rate_limit)])
async def write_extras(
    kind: Literal["graphical_abstract", "highlights"],
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    catalog = build_catalog(db, access.project, manuscript)
    ai = project_ai(db, access)
    run = new_ai_run(access, "manuscript", MANUSCRIPT_EXTRAS_PROMPT, ai)
    try:
        result = await write_extra(ai, kind, _review_summary(access.project), catalog)
    except LLMError as exc:
        run.status, run.error = "failed", str(exc)
        record_usage(run, ai, exc.usage)
        db.add(run)
        db.commit()
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    run.status = "succeeded"
    record_usage(run, ai, result.usage)
    db.add(run)
    db.flush()
    items = [i.strip() for i in result.value.items if i.strip()]
    manuscript.extras = {
        **manuscript.extras,
        kind: {
            "items": items,
            "unverified_numbers": unverified_numbers(items, catalog),
            "ai_run_id": run.id,
            "edited": False,
        },
    }
    _audit(db, access, f"ai.manuscript_{kind}", "manuscript", manuscript.id, {"items": items})
    db.commit()
    return manuscript.extras[kind]


@router.get("/manuscript/verification")
def manuscript_verification(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    return verify(db, access.project, _manuscript(db, access))


class AcknowledgementIn(BaseModel):
    section_key: str
    sentence_hash: str = Field(min_length=64, max_length=64)
    note: str = Field(min_length=10, max_length=2000)


@router.post("/manuscript/claims/acknowledgements", status_code=201)
def acknowledge_claim(
    body: AcknowledgementIn,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Acknowledge a sentence that needs no linked evidence (for example background with a citation-free statement)."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, body.section_key)
    catalog = build_catalog(db, access.project, manuscript)
    checks = {
        c.sentence.hash: c
        for c in check_content(section.content, catalog, reference_states(db, access.project.id), set())
    }
    check = checks.get(body.sentence_hash)
    if check is None:
        raise HTTPException(status_code=404, detail="That sentence isn't in the section")
    if check.status not in ACKNOWLEDGEABLE:
        raise HTTPException(
            status_code=409,
            detail="Only unsupported sentences, or numbers that can't be checked, can be acknowledged; fix the rest",
        )
    existing = db.scalar(
        select(models.ClaimAcknowledgement).where(
            models.ClaimAcknowledgement.section_id == section.id,
            models.ClaimAcknowledgement.sentence_hash == body.sentence_hash,
        )
    )
    if existing is None:
        existing = models.ClaimAcknowledgement(
            section_id=section.id,
            sentence_hash=body.sentence_hash,
            sentence=check.sentence.text,
            note=body.note,
            acknowledged_by_id=access.user.id,
        )
        db.add(existing)
        db.flush()
        _audit(
            db,
            access,
            "manuscript.claim_acknowledged",
            "manuscript_section",
            section.id,
            {"sentence": check.sentence.text, "status": check.status, "note": body.note},
        )
    db.commit()
    return verify(db, access.project, manuscript, catalog)


@router.delete("/manuscript/claims/acknowledgements")
def withdraw_acknowledgement(
    section_key: str,
    sentence_hash: str,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    section = _section(manuscript, section_key)
    row = db.scalar(
        select(models.ClaimAcknowledgement).where(
            models.ClaimAcknowledgement.section_id == section.id,
            models.ClaimAcknowledgement.sentence_hash == sentence_hash,
        )
    )
    if row is not None:
        _audit(
            db,
            access,
            "manuscript.claim_acknowledgement_withdrawn",
            "manuscript_section",
            section.id,
            {"sentence": row.sentence},
        )
        db.delete(row)
        db.commit()
    return verify(db, access.project, manuscript)


# --- References ---


@router.get("/manuscript/references")
def list_references(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    manuscript = _manuscript(db, access)
    refs = references(db, access.project.id)
    cited = cited_ids(manuscript)
    style = manuscript.citation_style
    rendered = citations.render(
        [s.content for s in sorted(manuscript.sections, key=lambda s: s.position)],
        {r.id: r.csl for r in refs},
        style if style in citations.BUILT_IN_STYLES else "vancouver",
    )
    numbers = {ref_id: number for number, ref_id, _ in rendered.bibliography}
    return {
        "style": style,
        "references": [{**reference_out(r, style, cited), "number": numbers.get(r.id)} for r in refs],
        "bibliography": [entry for _, _, entry in rendered.bibliography],
    }


class ReferenceIn(BaseModel):
    doi: str = Field("", max_length=255)
    pmid: str = Field("", max_length=20)
    record_id: int | None = None
    csl: dict | None = None


def _apply_check(ref: models.ManuscriptReference) -> None:
    try:
        result = check_reference(ref.csl, ref.doi, ref.pmid)
    except ReferenceCheckError as exc:
        ref.verification_status, ref.verification = "error", {"message": str(exc)}
        return
    ref.verification_status, ref.verification = result.status, result.details
    ref.retracted, ref.retraction = result.retracted, result.retraction
    ref.checked_at = models.utcnow()


@router.post("/manuscript/references", status_code=201)
async def add_reference(
    body: ReferenceIn,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    doi = body.doi.strip().lower().removeprefix("https://doi.org/")
    pmid = body.pmid.strip()
    existing = [
        r
        for r in references(db, access.project.id)
        if (doi and r.doi.lower() == doi) or (body.record_id and r.record_id == body.record_id)
    ]
    if existing:
        raise HTTPException(status_code=409, detail="This reference is already in the list")
    ref = models.ManuscriptReference(
        project_id=access.project.id, record_id=body.record_id, csl={}, doi=doi, pmid=pmid, created_by_id=access.user.id
    )
    db.add(ref)
    db.flush()
    try:
        if body.record_id is not None:
            record = get_in_project(db, models.Record, body.record_id, access.project.id, "Record")
            ref.csl = citations.csl_from_record(ref.id, record)
            ref.doi, ref.pmid = record.doi, (record.identifiers or {}).get("pmid", "")
        elif doi:
            work = await asyncio.to_thread(crossref_work, doi)
            if work is None:
                raise HTTPException(status_code=404, detail="Crossref has no record of this DOI")
            ref.csl = citations.csl_from_crossref(ref.id, work)
        elif pmid:
            summary = await asyncio.to_thread(pubmed_summary, pmid)
            if summary is None:
                raise HTTPException(status_code=404, detail="PubMed has no record of this PMID")
            ref.csl = citations.csl_from_pubmed(ref.id, summary)
            ref.doi = ref.csl.get("DOI", "")
        elif body.csl and body.csl.get("title"):
            ref.csl = {**body.csl, "id": str(ref.id)}
        else:
            raise HTTPException(status_code=422, detail="Give a DOI, PMID, record, or the reference's details")
    except ReferenceCheckError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    await asyncio.to_thread(_apply_check, ref)
    _audit(
        db,
        access,
        "manuscript.reference_added",
        "manuscript_reference",
        ref.id,
        {"doi": ref.doi, "pmid": ref.pmid, "verification_status": ref.verification_status, "retracted": ref.retracted},
    )
    db.commit()
    return reference_out(ref, manuscript.citation_style, cited_ids(manuscript))


class ReferenceUpdate(BaseModel):
    csl: dict


@router.put("/manuscript/references/{reference_id}")
def update_reference(
    reference_id: int,
    body: ReferenceUpdate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    ref = get_in_project(db, models.ManuscriptReference, reference_id, access.project.id, "Reference")
    if not body.csl.get("title"):
        raise HTTPException(status_code=422, detail="A reference needs a title")
    ref.csl = {**body.csl, "id": str(ref.id)}
    ref.doi = str(body.csl.get("DOI", ref.doi))
    ref.verification_status, ref.verification, ref.checked_at = "unchecked", {}, None
    _audit(db, access, "manuscript.reference_edited", "manuscript_reference", ref.id, {"csl": ref.csl})
    db.commit()
    return reference_out(ref, manuscript.citation_style, cited_ids(manuscript))


@router.delete("/manuscript/references/{reference_id}", status_code=204)
def delete_reference(
    reference_id: int,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    ref = get_in_project(db, models.ManuscriptReference, reference_id, access.project.id, "Reference")
    if ref.id in cited_ids(manuscript):
        raise HTTPException(status_code=409, detail="Remove the citations of this reference first")
    _audit(db, access, "manuscript.reference_deleted", "manuscript_reference", ref.id, {"doi": ref.doi})
    db.delete(ref)
    db.commit()
    return Response(status_code=204)


class CheckRequest(BaseModel):
    reference_ids: list[int] = Field(default_factory=list, max_length=500)


@router.post("/manuscript/references/check")
async def check_references(
    body: CheckRequest,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Verify DOIs and metadata against Crossref or PubMed, and look for retractions."""
    manuscript = _manuscript(db, access)
    refs = [r for r in references(db, access.project.id) if not body.reference_ids or r.id in body.reference_ids]
    for ref in refs:
        await asyncio.to_thread(_apply_check, ref)
    _audit(
        db,
        access,
        "manuscript.references_checked",
        "manuscript",
        manuscript.id,
        {
            "checked": len(refs),
            "retracted": [r.id for r in refs if r.retracted],
            "errors": [r.id for r in refs if r.verification_status == "error"],
        },
    )
    db.commit()
    return list_references(access, db)


@router.get("/manuscript/references/export")
def export_references(
    format: Literal["bibtex", "ris", "csljson"] = Query("bibtex"),
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    refs = {r.id: r.csl for r in references(db, access.project.id)}
    if format == "bibtex":
        return Response(
            citations.to_bibtex(refs),
            media_type="application/x-bibtex",
            headers={"Content-Disposition": 'attachment; filename="references.bib"'},
        )
    if format == "ris":
        return Response(
            citations.to_ris(refs),
            media_type="application/x-research-info-systems",
            headers={"Content-Disposition": 'attachment; filename="references.ris"'},
        )
    return Response(
        json.dumps(list(refs.values()), indent=2),
        media_type="application/json",
        headers={"Content-Disposition": 'attachment; filename="references.json"'},
    )


class ZoteroRequest(BaseModel):
    api_key: str = Field(min_length=10, max_length=200)
    library_type: Literal["user", "group"] = "user"
    library_id: str = Field(min_length=1, max_length=20)
    collection_key: str = Field("", max_length=20, pattern=r"^[A-Za-z0-9]*$")


@router.post("/manuscript/references/zotero/import")
async def import_from_zotero(
    body: ZoteroRequest,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Add references from a Zotero library or collection. The API key is used for this request only."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    try:
        items = await asyncio.to_thread(
            zotero_items, body.api_key, body.library_type, body.library_id, body.collection_key
        )
    except ReferenceCheckError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    existing = references(db, access.project.id)
    dois = {r.doi.lower() for r in existing if r.doi}
    titles = {str(r.csl.get("title", "")).casefold() for r in existing}
    added = 0
    for item in items:
        doi = str(item.get("DOI", "")).lower()
        title = str(item.get("title", "")).casefold()
        if not title or (doi and doi in dois) or title in titles:
            continue
        ref = models.ManuscriptReference(project_id=access.project.id, csl={}, doi=doi, created_by_id=access.user.id)
        db.add(ref)
        db.flush()
        ref.csl = {**item, "id": str(ref.id)}
        dois.add(doi)
        titles.add(title)
        added += 1
    _audit(
        db,
        access,
        "manuscript.references_imported_from_zotero",
        "manuscript",
        manuscript.id,
        {"items": len(items), "added": added},
    )
    db.commit()
    return {"items": len(items), "added": added}


@router.post("/manuscript/references/zotero/export")
async def export_to_zotero(
    body: ZoteroRequest,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    manuscript = _manuscript(db, access)
    items = [citations.zotero_item(r.csl) for r in references(db, access.project.id)]
    try:
        created = await asyncio.to_thread(
            zotero_create, body.api_key, body.library_type, body.library_id, items, body.collection_key
        )
    except ReferenceCheckError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    _audit(
        db,
        access,
        "manuscript.references_exported_to_zotero",
        "manuscript",
        manuscript.id,
        {"items": len(items), "created": created},
    )
    db.commit()
    return {"items": len(items), "created": created}


# --- Tables, figures, checklists ---


@router.get("/manuscript/assets")
def list_assets(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    return available_assets(db, access.project)


@router.get("/manuscript/tables/{key}")
def get_table(
    key: str, access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    item = table(db, access.project, key)
    if item is None:
        raise HTTPException(status_code=404, detail="Table not found")
    return {"key": item.key, "title": item.title, "header": item.header, "rows": item.rows, "note": item.note}


@router.get("/manuscript/figure")
async def get_figure(
    key: str,
    format: Literal["svg", "png", "pdf", "tiff"] = "svg",
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    item = await figure(db, access.project, key)
    if item is None or format not in item.files:
        raise HTTPException(status_code=404, detail="Figure not found in that format")
    media = {"svg": "image/svg+xml", "png": "image/png", "pdf": "application/pdf", "tiff": "image/tiff"}[format]
    return Response(
        item.files[format],
        media_type=media,
        headers={"Content-Disposition": f'inline; filename="{key.replace(":", "-")}.{format}"'},
    )


def _checklist_overrides(
    db: Session, manuscript: models.Manuscript, key: str
) -> dict[str, models.ManuscriptChecklistItem]:
    rows = db.scalars(
        select(models.ManuscriptChecklistItem).where(
            models.ManuscriptChecklistItem.manuscript_id == manuscript.id,
            models.ManuscriptChecklistItem.checklist == key,
        )
    )
    return {row.item_id: row for row in rows}


def _checklist_out(
    db: Session, project: models.Project, manuscript: models.Manuscript, key: str, context: dict
) -> dict:
    checklist = CHECKLISTS[key]
    labels = {s.key: s.label for s in PROTOCOL_SECTIONS}
    items = evaluate(checklist, manuscript, context, _checklist_overrides(db, manuscript, key), labels)
    return {
        "key": key,
        "label": checklist.label,
        "reference": checklist.reference,
        "url": checklist.url,
        "note": checklist.note,
        "applicable": applicable(checklist, context),
        "complete": complete(items),
        "items": items,
    }


@router.get("/manuscript/checklists")
def list_checklists(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    manuscript = _manuscript(db, access)
    context = checklist_context(db, access.project, manuscript)
    return [_checklist_out(db, access.project, manuscript, key, context) for key in CHECKLISTS]


class ChecklistItemIn(BaseModel):
    item_id: str = Field(max_length=20)
    status: str = Field("", max_length=20)
    location: str = Field("", max_length=2000)
    note: str = Field("", max_length=2000)


class ChecklistUpdate(BaseModel):
    items: list[ChecklistItemIn] = Field(min_length=1, max_length=100)


@router.put("/manuscript/checklists/{key}")
def update_checklist(
    key: str,
    body: ChecklistUpdate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    _open(db, access)
    manuscript = _manuscript(db, access)
    if key not in CHECKLISTS:
        raise HTTPException(status_code=404, detail="Checklist not found")
    valid = {item.id for item in CHECKLISTS[key].items}
    overrides = _checklist_overrides(db, manuscript, key)
    for entry in body.items:
        if entry.item_id not in valid:
            raise HTTPException(status_code=422, detail=f"{CHECKLISTS[key].label} has no item {entry.item_id}")
        if entry.status and entry.status not in STATUSES:
            raise HTTPException(status_code=422, detail=f"Unknown status: {entry.status}")
        row = overrides.get(entry.item_id) or models.ManuscriptChecklistItem(
            manuscript_id=manuscript.id, checklist=key, item_id=entry.item_id
        )
        row.status, row.location, row.note, row.updated_by_id = entry.status, entry.location, entry.note, access.user.id
        db.add(row)
    _audit(
        db,
        access,
        "manuscript.checklist_updated",
        "manuscript",
        manuscript.id,
        {"checklist": key, "items": {e.item_id: e.status for e in body.items}},
    )
    db.commit()
    return _checklist_out(db, access.project, manuscript, key, checklist_context(db, access.project, manuscript))


@router.get("/manuscript/checklists/{key}/export")
def export_checklist(
    key: str,
    format: Literal["csv", "docx"] = Query("docx"),
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    manuscript = _manuscript(db, access)
    if key not in CHECKLISTS:
        raise HTTPException(status_code=404, detail="Checklist not found")
    data = _checklist_out(db, access.project, manuscript, key, checklist_context(db, access.project, manuscript))
    header = ["Item", "Topic", "Status", "Location", "Note"]
    rows = [[i["item_id"], i["topic"], i["status"].replace("_", " "), i["location"], i["note"]] for i in data["items"]]
    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(header)
        writer.writerows(rows)
        return Response(
            buffer.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{key}-checklist.csv"'},
        )
    document = Document()
    document.add_heading(f"{data['label']} checklist: {manuscript.title}", level=1)
    document.add_paragraph(data["reference"])
    grid = document.add_table(rows=1, cols=len(header))
    grid.style = "Table Grid"
    for cell, label in zip(grid.rows[0].cells, header, strict=True):
        cell.text = label
    for values in rows:
        for cell, value in zip(grid.add_row().cells, values, strict=True):
            cell.text = value
    output = io.BytesIO()
    document.save(output)
    return Response(
        output.getvalue(),
        media_type=DOCX_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{key}-checklist.docx"'},
    )


# --- Authors, versions, approvals, export ---


class AuthorIn(BaseModel):
    user_id: int | None = None
    name: str = Field(min_length=1, max_length=300)
    email: str = Field("", max_length=320)
    affiliation: str = Field("", max_length=2000)
    orcid: str = Field("", max_length=40, pattern=r"^(|\d{4}-\d{4}-\d{4}-\d{3}[\dX])$")
    corresponding: bool = False
    credit_roles: list[str] = Field(default_factory=list, max_length=14)
    competing_interests: str = Field("", max_length=5000)


class AuthorsUpdate(BaseModel):
    authors: list[AuthorIn] = Field(max_length=100)


@router.put("/manuscript/authors")
def update_authors(
    body: AuthorsUpdate,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Set the author list. Changing it withdraws approvals of the current version."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    members = set(
        db.scalars(select(models.ProjectMember.user_id).where(models.ProjectMember.project_id == access.project.id))
    )
    if body.authors and sum(1 for a in body.authors if a.corresponding) != 1:
        raise HTTPException(status_code=422, detail="Mark exactly one corresponding author")
    for author in body.authors:
        if author.user_id is not None and author.user_id not in members:
            raise HTTPException(status_code=422, detail=f"{author.name} isn't a member of this project")
        unknown = set(author.credit_roles) - set(CREDIT_ROLES)
        if unknown:
            raise HTTPException(status_code=422, detail=f"Unknown CRediT roles: {', '.join(sorted(unknown))}")
    for existing in list(manuscript.authors):
        db.delete(existing)
    db.flush()
    manuscript.authors = [
        models.ManuscriptAuthor(position=i, **author.model_dump()) for i, author in enumerate(body.authors)
    ]
    _audit(
        db,
        access,
        "manuscript.authors_updated",
        "manuscript",
        manuscript.id,
        {"authors": [a.name for a in body.authors]},
    )
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


class VersionIn(BaseModel):
    note: str = Field("", max_length=2000)


@router.post("/manuscript/versions", status_code=201)
def create_version(
    body: VersionIn,
    access: ProjectAccess = Depends(project_access(Permission.EDIT_MANUSCRIPT)),
    db: Session = Depends(get_db),
):
    """Freeze the manuscript for author approval. Every sentence must be verified or acknowledged."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    if not manuscript.authors:
        raise HTTPException(status_code=409, detail="Add the authors first")
    report = verify(db, access.project, manuscript)
    if report["unresolved"]:
        raise HTTPException(
            status_code=409, detail=f"{report['unresolved']} sentences still need evidence, fixes, or acknowledgement"
        )
    cited = cited_ids(manuscript)
    retracted = [r.id for r in references(db, access.project.id) if r.id in cited and r.retracted]
    if retracted:
        raise HTTPException(status_code=409, detail="The manuscript cites retracted references")
    content = version_content(db, manuscript)
    sha = content_sha(content)
    latest = latest_version(manuscript)
    if latest is not None and latest.sha256 == sha:
        raise HTTPException(status_code=409, detail="Nothing changed since the latest version")
    version = models.ManuscriptVersion(
        manuscript_id=manuscript.id,
        number=(latest.number + 1) if latest else 1,
        content=content,
        sha256=sha,
        note=body.note,
        created_by_id=access.user.id,
    )
    db.add(version)
    db.flush()
    _audit(
        db,
        access,
        "manuscript.version_created",
        "manuscript_version",
        version.id,
        {"number": version.number, "sha256": sha},
    )
    db.commit()
    db.refresh(manuscript)
    return manuscript_out(db, access.project, manuscript)


@router.get("/manuscript/versions/{version_id}")
def get_version(
    version_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    manuscript = _manuscript(db, access)
    version = next((v for v in manuscript.versions if v.id == version_id), None)
    if version is None:
        raise HTTPException(status_code=404, detail="Version not found")
    return {
        "id": version.id,
        "number": version.number,
        "sha256": version.sha256,
        "note": version.note,
        "content": version.content,
        "created_at": version.created_at,
    }


@router.get("/manuscript/versions/{version_id}/compare/{other_id}")
def compare_versions(
    version_id: int,
    other_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    manuscript = _manuscript(db, access)
    by_id = {v.id: v for v in manuscript.versions}
    if version_id not in by_id or other_id not in by_id:
        raise HTTPException(status_code=404, detail="Version not found")
    before, after = by_id[version_id].content["sections"], by_id[other_id].content["sections"]
    return [
        {
            "key": key,
            "title": (after.get(key) or before.get(key))["title"],
            "diff": _word_diff((before.get(key) or {}).get("content", ""), (after.get(key) or {}).get("content", "")),
        }
        for key in dict.fromkeys([*before, *after])
        if (before.get(key) or {}).get("content") != (after.get(key) or {}).get("content")
    ]


class ApprovalIn(BaseModel):
    author_id: int | None = None
    note: str = Field("", max_length=2000)


@router.post("/manuscript/versions/{version_id}/approve")
def approve_version(
    version_id: int,
    body: ApprovalIn,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """An author approves the current version. The corresponding author may record the approval of a co-author without
    an OmniReview account, with a note saying how it was given."""
    _open(db, access)
    manuscript = _manuscript(db, access)
    latest = latest_version(manuscript)
    if latest is None or latest.id != version_id:
        raise HTTPException(status_code=409, detail="Only the latest version can be approved")
    if latest.sha256 != content_sha(version_content(db, manuscript)):
        raise HTTPException(status_code=409, detail="The manuscript changed after this version; create a new version")
    own = next((a for a in manuscript.authors if a.user_id == access.user.id), None)
    if body.author_id is None or (own is not None and body.author_id == own.id):
        if own is None:
            raise HTTPException(status_code=403, detail="Only the manuscript's authors can approve it")
        author = own
    else:
        other = next((a for a in manuscript.authors if a.id == body.author_id), None)
        if other is None:
            raise HTTPException(status_code=404, detail="Author not found")
        author = other
        if own is None or not own.corresponding or author.user_id is not None:
            raise HTTPException(
                status_code=403,
                detail=(
                    "Authors with an account approve for themselves; the corresponding author may record "
                    "approvals only for co-authors without one"
                ),
            )
        if len(body.note.strip()) < 10:
            raise HTTPException(
                status_code=422,
                detail="Say how the co-author gave their approval (for example by email, with the date)",
            )
    if any(a.author_id == author.id for a in latest.approvals):
        raise HTTPException(status_code=409, detail=f"{author.name} already approved this version")
    approval = models.AuthorApproval(
        version_id=latest.id, author_id=author.id, user_id=access.user.id, note=body.note.strip()
    )
    db.add(approval)
    db.flush()
    _audit(
        db,
        access,
        "manuscript.approved",
        "manuscript_version",
        latest.id,
        {
            "author": author.name,
            "recorded_for_other": author.user_id != access.user.id,
            "note": approval.note,
            "sha256": latest.sha256,
        },
    )
    db.commit()
    db.refresh(manuscript)
    return approval_state(db, manuscript)


@router.get("/manuscript/preview")
async def preview_manuscript(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    """The manuscript as Markdown with citations rendered and markers removed, for reading."""
    from manuscript_export import build_document

    document = await build_document(db, access.project, _manuscript(db, access), "md")
    return {"markdown": document.markdown}


@router.get("/manuscript/export")
async def export(
    format: Literal["docx", "tex", "pdf", "md"] = Query("docx"),
    access: ProjectAccess = Depends(project_access(Permission.EXPORT)),
    db: Session = Depends(get_db),
):
    """Export needs every author's approval of the current version and no retracted references (rechecked weekly)."""
    manuscript = _manuscript(db, access)
    state = approval_state(db, manuscript)
    if not state["all_approved"]:
        raise HTTPException(
            status_code=409, detail="Every author must approve the current manuscript version before export"
        )
    cited = cited_ids(manuscript)
    stale = [
        r
        for r in references(db, access.project.id)
        if r.id in cited and (r.checked_at is None or models.utcnow() - r.checked_at > RECHECK_AFTER)
    ]
    for ref in stale:
        await asyncio.to_thread(_apply_check, ref)
    db.commit()
    retracted = [r.id for r in references(db, access.project.id) if r.id in cited and r.retracted]
    if retracted:
        raise HTTPException(
            status_code=409, detail=f"Cited references have been retracted: {', '.join(str(i) for i in retracted)}"
        )
    try:
        content, media_type, filename = await export_manuscript(db, access.project, manuscript, format)
    except ExportUnavailable as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    _audit(
        db,
        access,
        "manuscript.exported",
        "manuscript",
        manuscript.id,
        {"format": format, "version_id": state["version_id"]},
    )
    db.commit()
    return Response(
        content, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def split_for_display(content: str) -> list[str]:
    return [s.text for s in split_sentences(content)]


__all__ = ["router", "sentence_hash", "split_for_display"]
