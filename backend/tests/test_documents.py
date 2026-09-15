"""Full texts: parsing into spans, uploads, open-access retrieval, background retrieval jobs, and PRISMA counts."""

import io
import os
from pathlib import Path

import pytest
import requests
from docx import Document as DocxFile
from workflow_helpers import add_member, create_project, decide, import_records, lock_protocol, open_screening, url

import documents
import documents_routes
import models
from database import SessionLocal
from document_parsing import SPAN_SEPARATOR, ParseError, detect_kind, parse_document
from services import fulltext

JATS = b"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Archiving and Interchange DTD v1.2 20190208//EN"
  "JATS-archivearticle1-mathml3.dtd">
<article xmlns:xlink="http://www.w3.org/1999/xlink" article-type="research-article">
<front><article-meta>
<title-group><article-title>Aspirin for <italic>primary</italic> prevention</article-title></title-group>
<permissions><license xlink:href="https://creativecommons.org/licenses/by/4.0/"><license-p>Open.</license-p></license>
</permissions>
<abstract><sec><title>Background</title><p>Aspirin may prevent heart attacks.</p></sec></abstract>
<abstract abstract-type="graphical"><p>Graphical abstract text.</p></abstract>
</article-meta></front>
<body>
<sec><title>Methods</title>
<p>We randomized 120 adults<xref ref-type="bibr" rid="r1">1</xref>.</p>
<sec><title>Outcomes</title><p>The primary outcome was myocardial infarction.
<table-wrap id="t1"><label>Table 1</label><caption><p>Baseline characteristics</p></caption>
<table><thead><tr><th>Group</th><th>n</th></tr></thead><tbody><tr><td>Aspirin</td><td>60</td></tr></tbody></table>
</table-wrap></p></sec>
</sec>
<sec><title>Results</title><p>Events were fewer with aspirin.</p>
<fig id="f1"><label>Figure 1</label><caption><p>Flow of participants</p></caption></fig></sec>
</body>
<back><ref-list><title>References</title>
<ref id="r1"><label>1</label><mixed-citation>Smith J. Aspirin trial. <source>BMJ</source>. 2001.</mixed-citation></ref>
</ref-list></back>
</article>"""

PAGES = [
    [
        "Aspirin for primary prevention",
        "",
        "Methods",
        "We randomized 120 adults to aspirin or placebo and followed them",
        "for five years in twelve general practices across the region.",
    ],
    [
        "Results",
        "Myocardial infarction occurred in 3 of 60 participants given aspirin",
        "and in 9 of 60 participants given placebo during follow-up.",
        "References",
        "1. Smith J. Aspirin trial. BMJ 2001.",
    ],
]


def make_pdf(pages: list[list[str]]) -> bytes:
    """A small text PDF with one line of Helvetica per string."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for index, lines in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * index} 0 R >>".encode()
        )
        operations = ["BT", "/F1 11 Tf", "14 TL", "72 720 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            operations.append(f"({escaped}) Tj T*")
        operations.append("ET")
        stream = "\n".join(operations).encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx() -> bytes:
    document = DocxFile()
    document.add_heading("Aspirin trial", level=0)
    document.add_heading("Methods", level=1)
    document.add_paragraph("We randomized 120 adults.")
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate([("Group", "n"), ("Aspirin", "60")]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Parsing ---


def test_jats_keeps_sections_tables_figures_and_references():
    parsed = parse_document(JATS, detect_kind("PMC1.xml", JATS))

    assert [(span.kind, span.section, span.label, span.text) for span in parsed.spans] == [
        ("title", "", "", "Aspirin for primary prevention"),
        ("heading", "Abstract > Background", "", "Background"),
        ("abstract", "Abstract > Background", "", "Aspirin may prevent heart attacks."),
        ("heading", "Methods", "", "Methods"),
        ("paragraph", "Methods", "", "We randomized 120 adults1."),
        ("heading", "Methods > Outcomes", "", "Outcomes"),
        ("paragraph", "Methods > Outcomes", "", "The primary outcome was myocardial infarction."),
        ("caption", "Methods > Outcomes", "Table 1", "Baseline characteristics"),
        ("table", "Methods > Outcomes", "Table 1", "Group | n\nAspirin | 60"),
        ("heading", "Results", "", "Results"),
        ("paragraph", "Results", "", "Events were fewer with aspirin."),
        ("caption", "Results", "Figure 1", "Flow of participants"),
        ("heading", "References", "", "References"),
        ("reference", "References", "1", "Smith J. Aspirin trial. BMJ. 2001."),
    ]
    assert fulltext.jats_license(JATS) == "cc-by"


def test_pdf_text_is_split_into_pages_and_sections():
    parsed = parse_document(make_pdf(PAGES), "pdf")

    assert parsed.page_count == 2
    assert parsed.parser.startswith("pypdf-")
    assert [(span.text, span.page) for span in parsed.spans if span.kind == "heading"] == [
        ("Methods", 1),
        ("Results", 2),
        ("References", 2),
    ]
    results = " ".join(span.text for span in parsed.spans if span.section == "Results")
    assert "3 of 60 participants given aspirin" in results
    assert any(span.kind == "reference" and "Smith J" in span.text for span in parsed.spans)


def test_scanned_pdfs_encrypted_or_broken_files_raise_readable_errors():
    with pytest.raises(ParseError, match="OCR"):
        parse_document(make_pdf([[""]]), "pdf")
    with pytest.raises(ParseError, match="couldn't be read"):
        parse_document(b"%PDF-1.4\nnot really a pdf", "pdf")
    with pytest.raises(ParseError, match="couldn't be read"):
        parse_document(b"<article><body>", "jats")


def test_docx_and_text_files_are_read():
    content = make_docx()
    parsed = parse_document(content, detect_kind("trial.docx", content))
    assert [(span.kind, span.section, span.text) for span in parsed.spans] == [
        ("title", "", "Aspirin trial"),
        ("heading", "Methods", "Methods"),
        ("paragraph", "Methods", "We randomized 120 adults."),
        ("table", "Methods", "Group | n\nAspirin | 60"),
    ]

    text = parse_document(b"Introduction\nAspirin\nworks.\n\nReferences\n1. Smith J.", detect_kind("notes.txt", b""))
    assert [(span.kind, span.text) for span in text.spans] == [
        ("heading", "Introduction"),
        ("paragraph", "Aspirin works."),
        ("heading", "References"),
        ("reference", "1. Smith J."),
    ]


def test_file_kinds_are_detected_from_content():
    assert detect_kind("a.bin", b"%PDF-1.7\n") == "pdf"
    assert detect_kind("page.xml", b"<!DOCTYPE html><html><article>x</article></html>") == "unsupported"
    assert detect_kind("data.csv", b"a,b\n1,2") == "unsupported"


# --- Retrieval ---


@pytest.mark.parametrize(
    "link",
    [
        "http://127.0.0.1/paper.pdf",
        "http://10.0.0.5/paper.pdf",
        "http://169.254.169.254/latest/meta-data",
        "http://[::1]/paper.pdf",
        "http://[::ffff:127.0.0.1]/paper.pdf",
        "ftp://example.org/paper.pdf",
        "https://user:secret@8.8.8.8/paper.pdf",
        "https://8.8.8.8:8080/paper.pdf",
    ],
)
def test_files_are_only_downloaded_from_public_addresses(link):
    with pytest.raises(fulltext.FetchError):
        fulltext.check_public_url(link)


def test_redirects_to_private_addresses_are_refused(monkeypatch):
    fulltext.check_public_url("https://8.8.8.8/paper.pdf")

    class Redirect:
        status_code = 302
        headers = {"location": "http://127.0.0.1/admin"}

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    requested = []

    def fake_get(link, **kwargs):
        requested.append(link)
        assert kwargs["allow_redirects"] is False
        return Redirect()

    monkeypatch.setattr(requests, "get", fake_get)
    with pytest.raises(fulltext.FetchError, match="private"):
        fulltext.download_public_file("https://8.8.8.8/paper.pdf", 1000)
    assert requested == ["https://8.8.8.8/paper.pdf"]


def test_europe_pmc_xml_is_preferred(monkeypatch):
    monkeypatch.setattr(fulltext, "europepmc_pmcid", lambda ids: "PMC42")
    monkeypatch.setattr(fulltext, "europepmc_full_text_xml", lambda pmcid, max_bytes: JATS)
    monkeypatch.setattr(fulltext, "unpaywall_lookup", lambda *args: pytest.fail("Unpaywall isn't needed"))

    result = fulltext.find_full_text(fulltext.record_ids("https://doi.org/10.1/A", {}), "me@example.org", 10_000)

    assert result.file is not None
    assert (result.file.origin, result.file.file_name, result.file.license) == ("europepmc", "PMC42.xml", "cc-by")
    assert [(a.source, a.outcome) for a in result.attempts] == [("europepmc", "found")]


def test_unpaywall_pdfs_are_tried_in_turn(monkeypatch):
    monkeypatch.setattr(fulltext, "europepmc_pmcid", lambda ids: "")
    monkeypatch.setattr(
        fulltext,
        "unpaywall_lookup",
        lambda doi, email: {
            "is_oa": True,
            "oa_status": "gold",
            "best_oa_location": {"url_for_pdf": "https://8.8.8.8/landing", "license": "cc-by"},
            "oa_locations": [
                {"url_for_pdf": "https://8.8.8.8/landing"},
                {"url_for_pdf": "https://8.8.4.4/paper.pdf", "license": "cc-by-nc", "version": "acceptedVersion"},
            ],
        },
    )
    pdf = make_pdf(PAGES)
    monkeypatch.setattr(
        fulltext,
        "download_public_file",
        lambda link, max_bytes: (b"<html>Sign in</html>", link) if "landing" in link else (pdf, link),
    )

    result = fulltext.find_full_text(fulltext.record_ids("10.1/a", {"pmid": "123"}), "me@example.org", 10**7)

    assert result.file is not None
    assert (result.file.origin, result.file.license, result.file.oa_status, result.file.version) == (
        "unpaywall",
        "cc-by-nc",
        "gold",
        "acceptedVersion",
    )
    assert [(a.source, a.outcome) for a in result.attempts] == [
        ("europepmc", "not_found"),
        ("unpaywall", "not_found"),
        ("unpaywall", "found"),
    ]


def test_records_without_identifiers_are_skipped():
    result = fulltext.find_full_text(fulltext.record_ids("", {"pmcid": "not-an-id"}), "", 1000)

    assert result.file is None
    assert [(a.source, a.outcome) for a in result.attempts] == [("europepmc", "skipped"), ("unpaywall", "skipped")]


# --- API ---


@pytest.fixture
def project(client, auth_headers):
    return create_project(client, auth_headers), auth_headers


@pytest.fixture
def screening(client, project, fake_provider):
    """A project with screening open and the first record (DOI 10.1/a) included."""
    project_id, headers = project
    first, second = open_screening(client, project_id, headers, fake_provider)
    decide(client, project_id, headers, first["id"], "include")
    return project_id, headers, first, second


def upload(client, project_id, headers, record_id, name, content, **form):
    return client.post(
        url(project_id, f"records/{record_id}/documents"), files={"file": (name, content)}, data=form, headers=headers
    )


def storage_path(document_id: int) -> Path:
    with SessionLocal() as db:
        document = db.get(models.Document, document_id)
        assert document is not None
        return Path(os.environ["DOCUMENT_STORAGE_DIR"]) / document.storage_key


def audit_actions(client, project_id, headers):
    return [event["action"] for event in client.get(url(project_id, "audit"), headers=headers).json()["events"]]


def test_uploaded_full_texts_are_parsed_downloadable_counted_and_deletable(client, screening):
    project_id, headers, first, _ = screening
    pdf = make_pdf(PAGES)

    response = upload(client, project_id, headers, first["id"], "trial.pdf", pdf)

    assert response.status_code == 201, response.text
    document = response.json()
    assert (document["origin"], document["role"], document["parse_status"], document["page_count"]) == (
        "upload",
        "full_text",
        "parsed",
        2,
    )
    detail = client.get(url(project_id, f"documents/{document['id']}"), headers=headers).json()
    assert detail["span_count"] == len(detail["spans"]) > 0
    plain_text = SPAN_SEPARATOR.join(span["text"] for span in detail["spans"])
    assert all(plain_text[span["start"] : span["end"]] == span["text"] for span in detail["spans"])
    download = client.get(url(project_id, f"documents/{document['id']}/file"), headers=headers)
    assert download.content == pdf
    assert download.headers["content-disposition"].startswith("attachment;")
    assert upload(client, project_id, headers, first["id"], "copy.pdf", pdf).status_code == 409

    listing = client.get(url(project_id, "full-texts"), headers=headers).json()
    assert listing["counts"] == {"sought": 1, "retrieved": 1, "not_retrieved": 0}
    assert [row["record"]["id"] for row in listing["records"]] == [first["id"]]
    prisma = client.get(url(project_id, "prisma"), headers=headers).json()
    assert (prisma["reports_sought_for_retrieval"], prisma["reports_not_retrieved"]) == (1, 0)

    stored = storage_path(document["id"])
    assert stored.read_bytes() == pdf
    assert client.delete(url(project_id, f"documents/{document['id']}"), headers=headers).status_code == 204
    assert not stored.exists()
    assert client.get(url(project_id, f"documents/{document['id']}"), headers=headers).status_code == 404
    assert client.get(url(project_id, "prisma"), headers=headers).json()["reports_not_retrieved"] == 1
    assert {"document.uploaded", "document.deleted"} <= set(audit_actions(client, project_id, headers))


def test_unreadable_files_are_kept_with_the_reason_and_can_be_reparsed(client, screening):
    project_id, headers, first, _ = screening

    scanned = upload(client, project_id, headers, first["id"], "scan.pdf", make_pdf([[""]])).json()
    supplement = upload(
        client, project_id, headers, first["id"], "data.xlsx", b"PK\x03\x04 spreadsheet", role="supplement"
    ).json()
    jats = upload(client, project_id, headers, first["id"], "PMC1.xml", JATS).json()

    assert (scanned["parse_status"], "OCR" in scanned["parse_error"]) == ("failed", True)
    assert (supplement["role"], supplement["parse_status"], supplement["media_type"]) == (
        "supplement",
        "unsupported",
        "application/octet-stream",
    )
    reparsed = client.post(url(project_id, f"documents/{jats['id']}/parse"), headers=headers)
    assert reparsed.status_code == 200, reparsed.text
    assert [span["kind"] for span in reparsed.json()["spans"]].count("table") == 1
    assert jats["span_count"] == len(reparsed.json()["spans"])


def test_documents_can_be_changed_only_during_screening_or_extraction(client, project, fake_provider):
    project_id, headers = project
    lock_protocol(client, project_id, headers, fake_provider)
    record = import_records(client, project_id, headers)[0]

    response = upload(client, project_id, headers, record["id"], "notes.txt", b"Methods\nRandomized.")

    assert response.status_code == 409
    assert "screening or extraction" in response.json()["detail"]


def test_uploads_are_size_limited_and_need_extraction_permission(client, screening, make_user, monkeypatch):
    project_id, headers, first, _ = screening
    viewer = add_member(client, project_id, headers, make_user, "viewer")

    assert upload(client, project_id, viewer, first["id"], "notes.txt", b"Methods").status_code == 403
    assert client.get(url(project_id, "full-texts"), headers=viewer).status_code == 200
    monkeypatch.setattr(documents_routes, "MAX_DOCUMENT_BYTES", 10)
    assert upload(client, project_id, headers, first["id"], "notes.txt", b"Methods\nRandomized.").status_code == 413


def fake_find_full_text(ids, email, max_bytes):
    """Europe PMC has the record with DOI 10.1/a; nothing else is found."""
    if ids.doi == "10.1/a":
        return fulltext.RetrievalResult(
            fulltext.FetchedFile(JATS, "PMC1.xml", "europepmc", "https://www.ebi.ac.uk/PMC1", license="cc-by"),
            [fulltext.Attempt("europepmc", "found", "JATS XML full text of PMC1")],
        )
    return fulltext.RetrievalResult(
        None,
        [
            fulltext.Attempt("europepmc", "skipped", "The record has no DOI, PMID, or PMCID"),
            fulltext.Attempt("unpaywall", "skipped", "Unpaywall needs a DOI"),
        ],
    )


def test_retrieving_one_record_stores_the_file_and_every_attempt(client, screening, monkeypatch):
    project_id, headers, first, second = screening
    monkeypatch.setattr(documents, "find_full_text", fake_find_full_text)

    found = client.post(url(project_id, f"records/{first['id']}/full-text/retrieve"), headers=headers)
    again = client.post(url(project_id, f"records/{first['id']}/full-text/retrieve"), headers=headers).json()
    missing = client.post(url(project_id, f"records/{second['id']}/full-text/retrieve"), headers=headers).json()

    assert found.status_code == 200, found.text
    body = found.json()
    assert body["retrieval"]["status"] == "found"
    assert (body["document"]["origin"], body["document"]["license"], body["document"]["parse_status"]) == (
        "europepmc",
        "cc-by",
        "parsed",
    )
    assert (again["retrieval"]["status"], again["document"]["id"]) == ("already_stored", body["document"]["id"])
    assert (missing["retrieval"]["status"], missing["document"]) == ("not_found", None)
    assert [attempt["outcome"] for attempt in missing["retrieval"]["attempts"]] == ["skipped", "skipped"]
    assert "document.retrieved" in audit_actions(client, project_id, headers)


def test_a_retrieval_job_covers_included_records_without_a_full_text(client, screening, monkeypatch):
    project_id, headers, first, second = screening
    decide(client, project_id, headers, second["id"], "include")
    monkeypatch.setattr(documents, "find_full_text", fake_find_full_text)

    job = client.post(url(project_id, "full-texts/retrieve"), json={}, headers=headers)

    assert job.status_code == 202, job.text
    assert (job.json()["task"], job.json()["status"], job.json()["processed"]) == ("fulltext", "completed", 2)
    listing = client.get(url(project_id, "full-texts"), headers=headers).json()
    assert listing["counts"] == {"sought": 2, "retrieved": 1, "not_retrieved": 1}
    latest = {row["record"]["id"]: row["latest_retrieval"]["status"] for row in listing["records"]}
    assert latest == {first["id"]: "found", second["id"]: "not_found"}
    assert "fulltext.retrieval_job" in audit_actions(client, project_id, headers)

    rerun = client.post(url(project_id, "full-texts/retrieve"), json={}, headers=headers).json()
    assert rerun["total"] == 1, "only the record still without a full text is retried"


def test_deleting_a_project_removes_its_files(client, screening):
    project_id, headers, first, _ = screening
    document = upload(client, project_id, headers, first["id"], "notes.txt", b"Methods\nRandomized adults.").json()
    stored = storage_path(document["id"])

    assert client.delete(f"/api/projects/{project_id}", headers=headers).status_code == 204

    assert not stored.exists()
    assert not stored.parent.exists()
