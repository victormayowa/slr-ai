"""Collaboration: invitations and ownership, competing-interest declarations, tasks, anchored comments with @mentions,
notifications, reviewer workload and metrics, and the team audit report.

Comments anchor to whatever they discuss (a record, a passage, an extraction cell, a manuscript sentence, a stage, or
the project) through an opaque anchor key, so a screen can show the thread beside the thing it's about.
"""

import hashlib
import io
import secrets
from collections import Counter, defaultdict
from datetime import date, timedelta
from typing import Any, Literal

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

import emailer
import entitlements
import help_center
import models
import notifications
from agreement import cohens_kappa
from audit import record_event
from auth_routes import get_current_user
from database import get_db
from permissions import Permission, ProjectRole
from projects_routes import (
    ProjectAccess,
    _ensure_not_last_owner,
    _find_member,
    _member_out,
    _require_owner,
    get_in_project,
    project_access,
)
from review_data import FULL_TEXT, TITLE_ABSTRACT
from workflow import STAGES

router = APIRouter(prefix="/api/projects/{project_id}", tags=["collaboration"])
invitations_router = APIRouter(prefix="/api/invitations", tags=["collaboration"])
notifications_router = APIRouter(prefix="/api/notifications", tags=["collaboration"])
me_router = APIRouter(prefix="/api/me", tags=["collaboration"])
help_router = APIRouter(prefix="/api/help", tags=["help"])

INVITATION_DAYS = 14
MAX_OPEN_INVITATIONS = 100


def _audit(db: Session, access: ProjectAccess, action: str, entity_type: str, entity_id: Any, details: dict) -> None:
    record_event(
        db,
        project_id=access.project.id,
        actor_id=access.user.id,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        details=details,
    )


# --- Invitations ---


class InvitationIn(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    role: ProjectRole


def invitation_out(invitation: models.Invitation) -> dict:
    status = "open"
    if invitation.accepted_at is not None:
        status = "accepted"
    elif invitation.revoked_at is not None:
        status = "revoked"
    elif invitation.expires_at <= models.utcnow():
        status = "expired"
    return {
        "id": invitation.id,
        "email": invitation.email,
        "role": invitation.role,
        "status": status,
        "invited_by": invitation.invited_by.full_name if invitation.invited_by else None,
        "expires_at": invitation.expires_at,
        "accepted_at": invitation.accepted_at,
        "created_at": invitation.created_at,
    }


@router.get("/invitations")
def list_invitations(
    access: ProjectAccess = Depends(project_access(Permission.MANAGE_MEMBERS)), db: Session = Depends(get_db)
):
    rows = db.scalars(
        select(models.Invitation)
        .where(models.Invitation.project_id == access.project.id)
        .options(selectinload(models.Invitation.invited_by))
        .order_by(models.Invitation.id.desc())
    )
    return [invitation_out(row) for row in rows]


@router.post("/invitations", status_code=201)
def create_invitation(
    body: InvitationIn,
    access: ProjectAccess = Depends(project_access(Permission.MANAGE_MEMBERS)),
    db: Session = Depends(get_db),
):
    """Invite someone by email. Only an owner can invite another owner, and the link is valid for 14 days."""
    if body.role == ProjectRole.OWNER:
        _require_owner(access)
    email = body.email.strip().lower()
    existing = db.scalar(
        select(models.User)
        .join(models.ProjectMember, models.ProjectMember.user_id == models.User.id)
        .where(models.ProjectMember.project_id == access.project.id, func.lower(models.User.email) == email)
    )
    if existing is not None:
        raise HTTPException(status_code=409, detail="That person is already a member of this project")
    entitlements.require_members(db, access.project)
    open_count = db.scalar(
        select(func.count())
        .select_from(models.Invitation)
        .where(
            models.Invitation.project_id == access.project.id,
            models.Invitation.accepted_at.is_(None),
            models.Invitation.revoked_at.is_(None),
        )
    )
    if (open_count or 0) >= MAX_OPEN_INVITATIONS:
        raise HTTPException(status_code=409, detail="There are too many open invitations for this project")
    raw_token = secrets.token_urlsafe(32)
    invitation = models.Invitation(
        project_id=access.project.id,
        email=email,
        role=body.role,
        token_hash=hashlib.sha256(raw_token.encode()).hexdigest(),
        invited_by_id=access.user.id,
        expires_at=models.utcnow() + timedelta(days=INVITATION_DAYS),
    )
    db.add(invitation)
    db.flush()
    _audit(db, access, "invitation.created", "invitation", invitation.id, {"email": email, "role": body.role})
    link = f"/invitations/{raw_token}"
    invited = db.scalar(select(models.User).where(func.lower(models.User.email) == email))
    if invited is not None:
        notifications.notify(
            db,
            user_id=invited.id,
            project_id=access.project.id,
            kind="invitation",
            title=f"{access.user.full_name} invited you to {access.project.title}",
            body=f"You were invited as {body.role.replace('_', ' ')}.",
            link=link,
        )
    db.commit()
    role = body.role.replace("_", " ")
    try:
        emailer.send(
            email,
            f"{access.user.full_name} invited you to a review on OmniReview",
            f'{access.user.full_name} invited you to join "{access.project.title}" as {role}.\n\n'
            f"Open this link within {INVITATION_DAYS} days to accept. If you don't have an OmniReview account yet, "
            f"create one with this email address first:\n\n{notifications.app_url(link)}",
        )
        emailed = emailer.delivers()
    except emailer.EmailError:
        emailed = False
    # The link is also returned once, so the inviter can send it another way; only its hash is stored.
    return {**invitation_out(invitation), "link": notifications.app_url(link), "emailed": emailed}


@router.delete("/invitations/{invitation_id}", status_code=204)
def revoke_invitation(
    invitation_id: int,
    access: ProjectAccess = Depends(project_access(Permission.MANAGE_MEMBERS)),
    db: Session = Depends(get_db),
):
    invitation = get_in_project(db, models.Invitation, invitation_id, access.project.id, "Invitation")
    if invitation.accepted_at is not None:
        raise HTTPException(status_code=409, detail="That invitation has already been accepted")
    if invitation.revoked_at is None:
        invitation.revoked_at = models.utcnow()
        _audit(db, access, "invitation.revoked", "invitation", invitation.id, {"email": invitation.email})
        db.commit()
    return Response(status_code=204)


class AcceptIn(BaseModel):
    token: str = Field(min_length=10, max_length=200)


@invitations_router.post("/accept")
def accept_invitation(body: AcceptIn, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Join the project the invitation names. The invited address must be one of the signed-in account's addresses."""
    token_hash = hashlib.sha256(body.token.strip().encode()).hexdigest()
    invitation = db.scalar(select(models.Invitation).where(models.Invitation.token_hash == token_hash))
    if invitation is None:
        raise HTTPException(status_code=404, detail="That invitation link isn't valid")
    if invitation.accepted_at is not None:
        raise HTTPException(status_code=409, detail="That invitation has already been used")
    if invitation.revoked_at is not None:
        raise HTTPException(status_code=409, detail="That invitation was withdrawn")
    if invitation.expires_at <= models.utcnow():
        raise HTTPException(status_code=409, detail="That invitation has expired; ask for a new one")
    addresses = {user.email.lower(), (user.institutional_email or "").lower()}
    if invitation.email.lower() not in addresses:
        raise HTTPException(
            status_code=403, detail=f"This invitation is for {invitation.email}. Sign in with that address."
        )
    membership = db.scalar(
        select(models.ProjectMember).where(
            models.ProjectMember.project_id == invitation.project_id, models.ProjectMember.user_id == user.id
        )
    )
    if membership is None:
        entitlements.require_members(db, invitation.project)
        membership = models.ProjectMember(project_id=invitation.project_id, user_id=user.id, role=invitation.role)
        db.add(membership)
    invitation.accepted_at, invitation.accepted_by_id = models.utcnow(), user.id
    db.flush()
    record_event(
        db,
        project_id=invitation.project_id,
        actor_id=user.id,
        action="invitation.accepted",
        entity_type="invitation",
        entity_id=invitation.id,
        details={"email": invitation.email, "role": invitation.role},
    )
    db.commit()
    return {"project_id": invitation.project_id, "role": membership.role, "title": invitation.project.title}


class OwnershipTransfer(BaseModel):
    user_id: int
    # The role the current owner keeps.
    keep_role: ProjectRole = ProjectRole.LEAD_REVIEWER


@router.post("/ownership")
def transfer_ownership(
    body: OwnershipTransfer,
    access: ProjectAccess = Depends(project_access(Permission.MANAGE_MEMBERS)),
    db: Session = Depends(get_db),
):
    """Hand the project to another member. Only an owner can, and they can't make themselves the last owner's
    replacement without keeping a role."""
    _require_owner(access)
    if body.keep_role == ProjectRole.OWNER:
        raise HTTPException(status_code=422, detail="Choose the role you'll keep after handing over ownership")
    member = _find_member(access.project, body.user_id)
    if member.user_id == access.user.id:
        raise HTTPException(status_code=422, detail="Choose another member to take ownership")
    member.role = ProjectRole.OWNER
    access.project.owner_id = member.user_id
    access.membership.role = body.keep_role
    _ensure_not_last_owner(access.project, access.membership)
    _audit(
        db,
        access,
        "project.ownership_transferred",
        "project",
        access.project.id,
        {"to_user_id": member.user_id, "kept_role": body.keep_role},
    )
    notifications.notify(
        db,
        user_id=member.user_id,
        project_id=access.project.id,
        kind="stage",
        title=f"You are now the owner of {access.project.title}",
        body=f"{access.user.full_name} handed over ownership.",
        link=f"/projects/{access.project.id}/team",
    )
    db.commit()
    return [_member_out(m) for m in access.project.members]


# --- Declarations ---


class DeclarationIn(BaseModel):
    has_competing_interests: bool
    statement: str = Field("", max_length=5000)
    funding: str = Field("", max_length=5000)


def declaration_out(declaration: models.MemberDeclaration, member: models.ProjectMember) -> dict:
    return {
        "user_id": declaration.user_id,
        "name": member.user.full_name,
        "role": member.role,
        "has_competing_interests": declaration.has_competing_interests,
        "statement": declaration.statement,
        "funding": declaration.funding,
        "updated_at": declaration.updated_at,
    }


@router.get("/declarations")
def list_declarations(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    rows = {
        row.user_id: row
        for row in db.scalars(
            select(models.MemberDeclaration).where(models.MemberDeclaration.project_id == access.project.id)
        )
    }
    return [
        (
            declaration_out(rows[member.user_id], member)
            if member.user_id in rows
            else {
                "user_id": member.user_id,
                "name": member.user.full_name,
                "role": member.role,
                "has_competing_interests": None,
                "statement": "",
                "funding": "",
                "updated_at": None,
            }
        )
        for member in access.project.members
    ]


@router.put("/declarations/me")
def save_declaration(
    body: DeclarationIn,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """Each member declares their own competing interests; nobody can declare on someone else's behalf."""
    if body.has_competing_interests and not body.statement.strip():
        raise HTTPException(status_code=422, detail="Describe the competing interest")
    declaration = db.scalar(
        select(models.MemberDeclaration).where(
            models.MemberDeclaration.project_id == access.project.id,
            models.MemberDeclaration.user_id == access.user.id,
        )
    )
    if declaration is None:
        declaration = models.MemberDeclaration(project_id=access.project.id, user_id=access.user.id)
        db.add(declaration)
    declaration.has_competing_interests = body.has_competing_interests
    declaration.statement = body.statement.strip()
    declaration.funding = body.funding.strip()
    db.flush()
    _audit(
        db,
        access,
        "declaration.saved",
        "declaration",
        declaration.id,
        {"has_competing_interests": body.has_competing_interests},
    )
    db.commit()
    return declaration_out(declaration, access.membership)


# --- Tasks ---


class TaskIn(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    description: str = Field("", max_length=5000)
    stage: str = Field("", max_length=30)
    assignee_id: int | None = None
    due_on: date | None = None
    priority: Literal["low", "normal", "high"] = "normal"
    anchor_key: str = Field("", max_length=200)


class TaskUpdate(BaseModel):
    title: str | None = Field(None, min_length=1, max_length=300)
    description: str | None = Field(None, max_length=5000)
    stage: str | None = Field(None, max_length=30)
    assignee_id: int | None = None
    due_on: date | None = None
    status: Literal["open", "in_progress", "done"] | None = None
    priority: Literal["low", "normal", "high"] | None = None


def task_out(task: models.Task) -> dict:
    return {
        "id": task.id,
        "title": task.title,
        "description": task.description,
        "stage": task.stage,
        "assignee_id": task.assignee_id,
        "assignee": task.assignee.full_name if task.assignee else None,
        "due_on": task.due_on,
        "status": task.status,
        "priority": task.priority,
        "anchor_key": task.anchor_key,
        "completed_at": task.completed_at,
        "created_at": task.created_at,
        "updated_at": task.updated_at,
    }


def _check_stage(stage: str) -> str:
    stage = stage.strip()
    if stage and stage not in STAGES:
        raise HTTPException(status_code=422, detail=f"Unknown workflow stage: {stage}")
    return stage


def _check_assignee(access: ProjectAccess, assignee_id: int | None) -> int | None:
    if assignee_id is None:
        return None
    _find_member(access.project, assignee_id)
    return assignee_id


def _may_change_task(access: ProjectAccess, task: models.Task) -> bool:
    return (
        task.created_by_id == access.user.id
        or task.assignee_id == access.user.id
        or access.can(Permission.MANAGE_WORKFLOW)
    )


@router.get("/tasks")
def list_tasks(
    status: Literal["open", "in_progress", "done", "all"] = "all",
    assignee_id: int | None = None,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    query = (
        select(models.Task)
        .where(models.Task.project_id == access.project.id)
        .options(selectinload(models.Task.assignee))
    )
    if status != "all":
        query = query.where(models.Task.status == status)
    if assignee_id is not None:
        query = query.where(models.Task.assignee_id == assignee_id)
    tasks = db.scalars(query.order_by(models.Task.status, models.Task.due_on.nulls_last(), models.Task.id))
    return [task_out(task) for task in tasks]


@router.post("/tasks", status_code=201)
def create_task(
    body: TaskIn,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    task = models.Task(
        project_id=access.project.id,
        title=body.title.strip(),
        description=body.description.strip(),
        stage=_check_stage(body.stage),
        assignee_id=_check_assignee(access, body.assignee_id),
        due_on=body.due_on,
        priority=body.priority,
        anchor_key=body.anchor_key.strip(),
        created_by_id=access.user.id,
    )
    db.add(task)
    db.flush()
    _audit(db, access, "task.created", "task", task.id, {"title": task.title, "assignee_id": task.assignee_id})
    if task.assignee_id and task.assignee_id != access.user.id:
        notifications.notify(
            db,
            user_id=task.assignee_id,
            project_id=access.project.id,
            kind="task_assigned",
            title=f"{access.user.full_name} assigned you: {task.title}",
            body=task.description[:500],
            link=f"/projects/{access.project.id}/team",
        )
    db.commit()
    return task_out(task)


@router.patch("/tasks/{task_id}")
def update_task(
    task_id: int,
    body: TaskUpdate,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    task = get_in_project(db, models.Task, task_id, access.project.id, "Task")
    if not _may_change_task(access, task):
        raise HTTPException(status_code=403, detail="Only the task's author, its assignee, or a lead can change it")
    previous_assignee = task.assignee_id
    fields = body.model_dump(exclude_unset=True)
    if "stage" in fields and fields["stage"] is not None:
        fields["stage"] = _check_stage(fields["stage"])
    if "assignee_id" in fields:
        fields["assignee_id"] = _check_assignee(access, fields["assignee_id"])
    for name, value in fields.items():
        setattr(task, name, value)
    if body.status is not None:
        task.completed_at = models.utcnow() if body.status == "done" else None
    if "due_on" in fields:
        task.due_reminder_sent_at = None
    db.flush()
    _audit(db, access, "task.updated", "task", task.id, {"changed": sorted(fields), "status": task.status})
    if task.assignee_id and task.assignee_id != previous_assignee and task.assignee_id != access.user.id:
        notifications.notify(
            db,
            user_id=task.assignee_id,
            project_id=access.project.id,
            kind="task_assigned",
            title=f"{access.user.full_name} assigned you: {task.title}",
            body=task.description[:500],
            link=f"/projects/{access.project.id}/team",
        )
    db.commit()
    return task_out(task)


@router.delete("/tasks/{task_id}", status_code=204)
def delete_task(
    task_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    task = get_in_project(db, models.Task, task_id, access.project.id, "Task")
    if not _may_change_task(access, task):
        raise HTTPException(status_code=403, detail="Only the task's author, its assignee, or a lead can delete it")
    _audit(db, access, "task.deleted", "task", task.id, {"title": task.title})
    db.delete(task)
    db.commit()
    return Response(status_code=204)


# --- Comments ---


class CommentIn(BaseModel):
    anchor_key: str = Field(min_length=1, max_length=200)
    anchor_label: str = Field("", max_length=500)
    body: str = Field(min_length=1, max_length=10_000)
    parent_id: int | None = None


class CommentUpdate(BaseModel):
    body: str = Field(min_length=1, max_length=10_000)


def comment_out(comment: models.Comment) -> dict:
    return {
        "id": comment.id,
        "anchor_key": comment.anchor_key,
        "anchor_label": comment.anchor_label,
        "parent_id": comment.parent_id,
        "body": "" if comment.deleted else comment.body,
        "deleted": comment.deleted,
        "mentions": comment.mentions,
        "author_id": comment.author_id,
        "author": comment.author.full_name if comment.author else None,
        "resolved": comment.resolved,
        "resolved_at": comment.resolved_at,
        "edited_at": comment.edited_at,
        "created_at": comment.created_at,
    }


@router.get("/comments")
def list_comments(
    anchor_key: str | None = Query(None, max_length=200),
    anchor_prefix: str | None = Query(None, max_length=200),
    unresolved_only: bool = False,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    query = (
        select(models.Comment)
        .where(models.Comment.project_id == access.project.id)
        .options(selectinload(models.Comment.author))
    )
    if anchor_key:
        query = query.where(models.Comment.anchor_key == anchor_key)
    if anchor_prefix:
        query = query.where(models.Comment.anchor_key.startswith(anchor_prefix))
    if unresolved_only:
        query = query.where(models.Comment.resolved.is_(False))
    return [comment_out(row) for row in db.scalars(query.order_by(models.Comment.id))]


@router.get("/comments/counts")
def comment_counts(
    anchor_prefix: str = Query("", max_length=200),
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """How many comments and unresolved threads each anchor has, so screens can show a badge without loading them."""
    query = select(
        models.Comment.anchor_key,
        func.count(),
        func.count().filter(models.Comment.resolved.is_(False)),
    ).where(models.Comment.project_id == access.project.id, models.Comment.deleted.is_(False))
    if anchor_prefix:
        query = query.where(models.Comment.anchor_key.startswith(anchor_prefix))
    rows = db.execute(query.group_by(models.Comment.anchor_key)).all()
    return {key: {"comments": total, "unresolved": unresolved} for key, total, unresolved in rows}


@router.post("/comments", status_code=201)
def create_comment(
    body: CommentIn,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """Add a comment. Replies take their thread's anchor, and @email mentions notify that member."""
    anchor_key, anchor_label = body.anchor_key.strip(), body.anchor_label.strip()
    parent = None
    if body.parent_id is not None:
        parent = get_in_project(db, models.Comment, body.parent_id, access.project.id, "Comment")
        if parent.parent_id is not None:
            raise HTTPException(status_code=422, detail="Reply to the first comment of the thread")
        anchor_key, anchor_label = parent.anchor_key, parent.anchor_label
    mentions = notifications.mentioned_users(body.body, access.project.members)
    comment = models.Comment(
        project_id=access.project.id,
        anchor_key=anchor_key,
        anchor_label=anchor_label,
        parent_id=body.parent_id,
        body=body.body.strip(),
        mentions=mentions,
        author_id=access.user.id,
    )
    db.add(comment)
    db.flush()
    link = f"/projects/{access.project.id}/team?comment={comment.id}"
    for user_id in mentions:
        if user_id != access.user.id:
            notifications.notify(
                db,
                user_id=user_id,
                project_id=access.project.id,
                kind="mention",
                title=f"{access.user.full_name} mentioned you",
                body=comment.body[:500],
                link=link,
            )
    if parent is not None:
        thread = db.scalars(
            select(models.Comment).where((models.Comment.id == parent.id) | (models.Comment.parent_id == parent.id))
        ).all()
        already_told = {access.user.id, *mentions}
        for user_id in sorted({c.author_id for c in thread if c.author_id is not None} - already_told):
            notifications.notify(
                db,
                user_id=user_id,
                project_id=access.project.id,
                kind="reply",
                title=f"{access.user.full_name} replied to your comment",
                body=comment.body[:500],
                link=link,
            )
    db.commit()
    return comment_out(comment)


@router.patch("/comments/{comment_id}")
def update_comment(
    comment_id: int,
    body: CommentUpdate,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    comment = get_in_project(db, models.Comment, comment_id, access.project.id, "Comment")
    if comment.author_id != access.user.id:
        raise HTTPException(status_code=403, detail="Only the author can edit a comment")
    if comment.deleted:
        raise HTTPException(status_code=409, detail="That comment was deleted")
    comment.body = body.body.strip()
    comment.mentions = notifications.mentioned_users(comment.body, access.project.members)
    comment.edited_at = models.utcnow()
    db.commit()
    return comment_out(comment)


@router.delete("/comments/{comment_id}", status_code=204)
def delete_comment(
    comment_id: int,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """Comments are kept as a deleted placeholder so replies keep their thread."""
    comment = get_in_project(db, models.Comment, comment_id, access.project.id, "Comment")
    if comment.author_id != access.user.id and not access.can(Permission.MANAGE_WORKFLOW):
        raise HTTPException(status_code=403, detail="Only the author or a lead can delete a comment")
    comment.deleted, comment.body, comment.mentions = True, "", []
    db.commit()
    return Response(status_code=204)


@router.post("/comments/{comment_id}/resolve")
def resolve_comment(
    comment_id: int,
    resolved: bool = True,
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)),
    db: Session = Depends(get_db),
):
    """Resolve or reopen a thread. Resolving the first comment resolves its replies."""
    comment = get_in_project(db, models.Comment, comment_id, access.project.id, "Comment")
    thread = [comment] + list(db.scalars(select(models.Comment).where(models.Comment.parent_id == comment.id)))
    for item in thread:
        item.resolved = resolved
        item.resolved_by_id = access.user.id if resolved else None
        item.resolved_at = models.utcnow() if resolved else None
    db.commit()
    return comment_out(comment)


# --- Notifications ---


def notification_out(notification: models.Notification) -> dict:
    return {
        "id": notification.id,
        "project_id": notification.project_id,
        "kind": notification.kind,
        "title": notification.title,
        "body": notification.body,
        "link": notification.link,
        "read_at": notification.read_at,
        "created_at": notification.created_at,
    }


@notifications_router.get("")
def list_notifications(
    unread_only: bool = False,
    limit: int = Query(50, ge=1, le=200),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = select(models.Notification).where(models.Notification.user_id == user.id)
    if unread_only:
        query = query.where(models.Notification.read_at.is_(None))
    rows = db.scalars(query.order_by(models.Notification.id.desc()).limit(limit)).all()
    unread = db.scalar(
        select(func.count())
        .select_from(models.Notification)
        .where(models.Notification.user_id == user.id, models.Notification.read_at.is_(None))
    )
    return {"unread": unread or 0, "notifications": [notification_out(row) for row in rows]}


@notifications_router.post("/{notification_id}/read")
def mark_notification_read(
    notification_id: int, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)
):
    notification = db.get(models.Notification, notification_id)
    if notification is None or notification.user_id != user.id:
        raise HTTPException(status_code=404, detail="Notification not found")
    if notification.read_at is None:
        notification.read_at = models.utcnow()
        db.commit()
    return notification_out(notification)


@notifications_router.post("/read-all")
def mark_all_read(user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.scalars(
        select(models.Notification).where(models.Notification.user_id == user.id, models.Notification.read_at.is_(None))
    ).all()
    for notification in rows:
        notification.read_at = models.utcnow()
    db.commit()
    return {"marked": len(rows)}


class EmailPreference(BaseModel):
    email_notifications: bool


@me_router.get("/notification-preferences")
def get_notification_preferences(user: models.User = Depends(get_current_user)):
    return {"email_notifications": user.email_notifications, "email_configured": notifications.smtp_configured()}


@me_router.put("/notification-preferences")
def set_notification_preferences(
    body: EmailPreference, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)
):
    user.email_notifications = body.email_notifications
    db.commit()
    return {"email_notifications": user.email_notifications, "email_configured": notifications.smtp_configured()}


# --- Workload and reviewer metrics ---


@router.get("/team/workload")
def team_workload(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    """Each member's screening, extraction, and appraisal counts, and their open tasks."""
    decisions: Counter[tuple[int, str]] = Counter(
        (reviewer_id, stage)
        for reviewer_id, stage in db.execute(
            select(models.ScreeningDecision.reviewer_id, models.ScreeningDecision.stage)
            .join(models.Record, models.Record.id == models.ScreeningDecision.record_id)
            .where(models.Record.project_id == access.project.id)
        )
    )
    extractions = Counter(
        db.scalars(
            select(models.ExtractionValue.extractor_id).where(models.ExtractionValue.project_id == access.project.id)
        ).all()
    )
    appraisals = Counter(
        db.scalars(
            select(models.AppraisalAssessment.signed_off_by_id).where(
                models.AppraisalAssessment.project_id == access.project.id,
                models.AppraisalAssessment.status == "signed_off",
            )
        ).all()
    )
    tasks = Counter(
        db.scalars(
            select(models.Task.assignee_id).where(
                models.Task.project_id == access.project.id, models.Task.status != "done"
            )
        ).all()
    )
    overdue = Counter(
        db.scalars(
            select(models.Task.assignee_id).where(
                models.Task.project_id == access.project.id,
                models.Task.status != "done",
                models.Task.due_on.is_not(None),
                models.Task.due_on < models.utcnow().date(),
            )
        ).all()
    )
    return [
        {
            "user_id": member.user_id,
            "name": member.user.full_name,
            "role": member.role,
            "title_abstract_decisions": decisions[(member.user_id, TITLE_ABSTRACT)],
            "full_text_decisions": decisions[(member.user_id, FULL_TEXT)],
            "extraction_values": extractions[member.user_id],
            "appraisals_signed_off": appraisals[member.user_id],
            "open_tasks": tasks[member.user_id],
            "overdue_tasks": overdue[member.user_id],
        }
        for member in access.project.members
    ]


@router.get("/team/metrics")
def reviewer_metrics(
    access: ProjectAccess = Depends(project_access(Permission.VIEW_PROJECT)), db: Session = Depends(get_db)
):
    """Per reviewer: decisions, how often they agreed with the final decision, and how often they differed from the AI
    suggestion. Agreement between each pair of reviewers is reported as Cohen's kappa."""
    records = db.scalars(
        select(models.Record)
        .where(models.Record.project_id == access.project.id, models.Record.duplicate_of_id.is_(None))
        .options(
            selectinload(models.Record.decisions),
            selectinload(models.Record.adjudications),
            selectinload(models.Record.ai_runs).selectinload(models.AIRun.screening),
        )
    ).all()
    by_reviewer: dict[int, dict[str, int]] = defaultdict(
        lambda: {"decisions": 0, "agreed_with_final": 0, "compared_with_final": 0, "differed_from_ai": 0, "ai_seen": 0}
    )
    pairs: dict[tuple[int, int], list[tuple[str, str]]] = defaultdict(list)
    for record in records:
        suggestion = next(
            (
                run.screening
                for run in reversed(record.ai_runs)
                if run.screening and run.screening.stage == TITLE_ABSTRACT
            ),
            None,
        )
        ai_decision = None
        if suggestion is not None:
            ai_decision = "include" if suggestion.decision in ("Include", "Maybe") else "exclude"
        for stage in (TITLE_ABSTRACT, FULL_TEXT):
            stage_decisions = [d for d in record.decisions if d.stage == stage and d.decision != "undecided"]
            adjudication = next((a for a in record.adjudications if a.stage == stage), None)
            counts = Counter(d.decision for d in stage_decisions)
            final = (
                adjudication.decision if adjudication else (counts.most_common(1)[0][0] if len(counts) == 1 else None)
            )
            for decision in stage_decisions:
                stats = by_reviewer[decision.reviewer_id]
                stats["decisions"] += 1
                if final is not None:
                    stats["compared_with_final"] += 1
                    stats["agreed_with_final"] += int(decision.decision == final)
                if stage == TITLE_ABSTRACT and ai_decision is not None:
                    stats["ai_seen"] += 1
                    stats["differed_from_ai"] += int(decision.decision != ai_decision)
            if stage == TITLE_ABSTRACT:
                for first, second in (
                    (a, b) for a in stage_decisions for b in stage_decisions if a.reviewer_id < b.reviewer_id
                ):
                    pairs[(first.reviewer_id, second.reviewer_id)].append((first.decision, second.decision))
    names = {member.user_id: member.user.full_name for member in access.project.members}
    reviewers = [
        {
            "user_id": user_id,
            "name": names.get(user_id, "Former member"),
            **stats,
            "agreement_with_final": (
                stats["agreed_with_final"] / stats["compared_with_final"] if stats["compared_with_final"] else None
            ),
            "override_rate": stats["differed_from_ai"] / stats["ai_seen"] if stats["ai_seen"] else None,
        }
        for user_id, stats in sorted(by_reviewer.items())
    ]
    agreement = []
    for (first_id, second_id), items in sorted(pairs.items()):
        result = cohens_kappa(items, ["include", "exclude"])
        agreement.append(
            {
                "reviewers": [names.get(first_id, "Former member"), names.get(second_id, "Former member")],
                "n": result.n,
                "observed": result.observed,
                "kappa": result.kappa,
                "pabak": result.pabak,
            }
        )
    return {"reviewers": reviewers, "pairwise_agreement": agreement}


@router.get("/team/report")
def team_audit_report(
    format: Literal["json", "docx"] = "json",
    access: ProjectAccess = Depends(project_access(Permission.VIEW_AUDIT)),
    db: Session = Depends(get_db),
):
    """Who did what on the review: members and their roles, declarations, workload, agreement, and task completion."""
    workload = team_workload(access, db)
    metrics = reviewer_metrics(access, db)
    declarations = list_declarations(access, db)
    tasks = db.scalars(select(models.Task).where(models.Task.project_id == access.project.id)).all()
    events = db.scalar(
        select(func.count()).select_from(models.AuditEvent).where(models.AuditEvent.project_id == access.project.id)
    )
    report = {
        "project": {"id": access.project.id, "title": access.project.title},
        "generated_at": models.utcnow(),
        "members": [_member_out(member) for member in access.project.members],
        "declarations": declarations,
        "workload": workload,
        "metrics": metrics,
        "tasks": {
            "total": len(tasks),
            "done": sum(1 for task in tasks if task.status == "done"),
            "open": sum(1 for task in tasks if task.status != "done"),
            "overdue": sum(
                1
                for task in tasks
                if task.status != "done" and task.due_on is not None and task.due_on < models.utcnow().date()
            ),
        },
        "audit_events": events or 0,
    }
    if format == "json":
        return report
    document = WordDocument()
    document.add_heading(f"Team audit report: {access.project.title}", level=1)
    document.add_paragraph(f"Generated {report['generated_at']:%Y-%m-%d %H:%M} UTC")
    document.add_heading("Members and declarations", level=2)
    for declaration in declarations:
        competing = declaration["has_competing_interests"]
        statement = (
            "not yet declared"
            if competing is None
            else (declaration["statement"] or "declared")
            if competing
            else "none declared"
        )
        document.add_paragraph(
            f"{declaration['name']} ({declaration['role'].replace('_', ' ')}): competing interests {statement}"
        )
    document.add_heading("Workload", level=2)
    for row in workload:
        document.add_paragraph(
            f"{row['name']}: {row['title_abstract_decisions']} title and abstract decisions, "
            f"{row['full_text_decisions']} full-text decisions, {row['appraisals_signed_off']} appraisals signed off, "
            f"{row['open_tasks']} open tasks"
        )
    document.add_heading("Agreement between reviewers", level=2)
    if metrics["pairwise_agreement"]:
        for row in metrics["pairwise_agreement"]:
            kappa = "not estimable" if row["kappa"] is None else f"{row['kappa']:.2f}"
            document.add_paragraph(f"{' and '.join(row['reviewers'])}: kappa {kappa} over {row['n']} records")
    else:
        document.add_paragraph("Only one reviewer screened each record, so agreement isn't estimated.")
    document.add_heading("Tasks and audit trail", level=2)
    document.add_paragraph(
        f"{report['tasks']['done']} of {report['tasks']['total']} tasks done, "
        f"{report['tasks']['overdue']} overdue. {report['audit_events']} audit events recorded."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    name = f"team-report-project-{access.project.id}.docx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# --- Help ---


@help_router.get("/search")
def search_help(q: str = Query(min_length=2, max_length=300), user: models.User = Depends(get_current_user)):
    sections = help_center.search(q)
    return [{**help_center.section_out(section), "text": section.text[:1500]} for section in sections]
